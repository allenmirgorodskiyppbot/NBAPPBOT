import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from nbabot.core import Store,load_config
from nbabot.master_source import import_master,master_status,master_path,refresh_master,load_master_reference


class MasterSource(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.root=Path(self.tmp.name)
        self.config=load_config(self.root/'none')
        self.config.update(data_dir=str(self.root/'state'),data_source_mode='spreadsheet',
            nba_enabled=False,public_context_enabled=False,organized_players=[],simulations=100)
        self.store=Store(self.config)

    def tearDown(self):
        self.store.close();self.tmp.cleanup()

    def row(self,**kw):
        row={'game_id':'0022400001','player_id':'00123','name':'Test Guard','team':'NYK','opponent':'BOS',
            'season':'2024-25','season_type':'Regular Season','date_et':'2025-01-01','appeared':1,'started':None,
            'minutes':34.098333333333336,'min_seconds':2045.9,'minutes_exact_text':'34.098333333333336',
            'min_seconds_exact_text':'2045.9','source_minutes':'34:06','pts':10,'reb':0,'ast':None,
            'fgm':5,'fga':11,'fg3m':0,'fg3a':1,'ftm':0,'fta':0,'oreb':0,'dreb':0,
            'known_at':'2025-02-01T00:00:00+00:00','source':'https://example.test/never-fetch'}
        row.update(kw);return row

    def organized(self,rows=None,context=None,extra_tables=False):
        path=self.root/('master'+str(len(list(self.root.glob('master*.xlsx'))))+'.xlsx')
        book=Workbook();sheet=book.active;sheet.title='Data_PlayerGames'
        sheet.append(['A view title']);sheet.append(['Source note']);sheet.append([])
        records=rows or [self.row()]
        headers=list(records[0]);sheet.append(headers)
        for row in records:sheet.append([row.get(k) for k in headers])
        from openpyxl.utils import get_column_letter
        sheet.add_table(Table(displayName='CanonicalGames',ref='A4:'+get_column_letter(len(headers))+str(4+len(records))))
        if context:
            sheet=book.create_sheet('Data_GameContext');sheet.append(list(context));sheet.append(list(context.values()))
            sheet.add_table(Table(displayName='Schedule',ref='A1:'+get_column_letter(len(context))+'2'))
            sheet.append([]);sheet.append(['game_id','player_id','started']);sheet.append(['wrong-role-game','00123',1])
            sheet.add_table(Table(displayName='ObservedRoles',ref='A4:C5'))
        if extra_tables:
            sheet=book.create_sheet('NYK_vs_BOS')
            h=['knicks_id','knicks_name','opponent_id','opponent_name','opponent_team','assignment','assignment_status','scope_note']
            sheet.append(h);sheet.append(['00123','Test Guard','00456','Named Opponent','BOS','primary','Estimated','Not a measured assignment'])
            sheet.add_table(Table(displayName='ExactPairs',ref='A1:H2'))
        book.save(path);return path

    def test_organized_ids_exact_minutes_missing_zero_and_source_time(self):
        path=self.organized(extra_tables=True);original=path.read_bytes()
        result=import_master(self.store,path)
        self.assertEqual(result['status'],'success')
        row=self.store.one('SELECT * FROM player_games')
        self.assertEqual(row['game_id'],'0022400001');self.assertEqual(row['player_id'],'00123')
        self.assertEqual(row['minutes'],34.098333333333336);self.assertEqual(row['min_seconds'],2045.9)
        self.assertEqual(row['reb'],0);self.assertIsNone(row['ast'])
        self.assertEqual(row['known_at'],'2025-02-01T00:00:00.000000+00:00')
        self.assertEqual(self.store.one("SELECT name FROM players WHERE id='00123'")['name'],'Test Guard')
        self.assertEqual(self.store.one('SELECT count(*) n FROM freshness')['n'],0)
        self.assertEqual(master_path(self.store).read_bytes(),original)
        self.assertEqual(path.read_bytes(),original)
        pair=json.loads(self.store.one("SELECT payload FROM raw_rows WHERE sheet='Matchups'")['payload'])
        self.assertEqual(pair['Suspected defender'],'Named Opponent')
        self.assertEqual(self.store.one("SELECT name FROM players WHERE id='00456'")['name'],'Named Opponent')

    def test_original_clean_stats_and_explicit_registry_utc(self):
        path=self.root/'original.xlsx';book=Workbook();sheet=book.active;sheet.title='Clean Stats'
        sheet.append(['Title']);sheet.append([])
        headers=['NBA game ID','NBA player ID','Player','Team','Opponent','Season','Game date','Appeared','Minutes (source precision)','PTS','REB','AST','FGM','FGA','3PM','3PA','FTM','FTA','Source ID']
        sheet.append(headers);sheet.append(['0022400001','00123','Test Guard','NYK','BOS','2024-25','2025-01-01','Yes',34.098333333333336,10,0,None,5,11,0,1,0,0,'source-one'])
        sheet=book.create_sheet('Source Registry');sheet.append(['Source ID','Source URL','Retrieved at (UTC)'])
        sheet.append(['source-one','https://example.test/no-network','2025-02-01 00:00:00'])
        book.save(path)
        import_master(self.store,path)
        self.assertEqual(self.store.one('SELECT known_at FROM player_games')['known_at'],'2025-02-01T00:00:00.000000+00:00')

    def test_identical_duplicate_dedup_and_playoffs_separate(self):
        path=self.organized([self.row(),self.row(),self.row(game_id='0042400001',season='2024-25',season_type='Playoffs',date_et='2025-05-01')])
        result=import_master(self.store,path)
        self.assertEqual(result['duplicates_removed'],1)
        self.assertEqual(self.store.one('SELECT count(*) n FROM player_games')['n'],1)
        self.assertEqual(master_status(self.store)['coverage']['playoff_games'],1)
        self.assertEqual(len(load_master_reference(self.store)['playoff_games']),1)

    def test_bad_duplicate_or_join_retains_old_database_and_pointer(self):
        import_master(self.store,self.organized())
        before=master_status(self.store)['file_hash']
        bad=self.organized([self.row(),self.row(pts=12,fgm=6)])
        with self.assertRaisesRegex(ValueError,'Conflicting duplicate'):import_master(self.store,bad)
        bad=self.organized([self.row(),self.row(player_id='00456',date_et='2025-01-02')])
        with self.assertRaisesRegex(ValueError,'different dates'):import_master(self.store,bad)
        self.assertEqual(master_status(self.store)['file_hash'],before)
        self.assertEqual(self.store.one('SELECT pts FROM player_games')['pts'],10)

    def test_failed_transaction_retains_previous_records_and_active_original(self):
        import_master(self.store,self.organized())
        before=master_status(self.store)
        new=self.organized([self.row(pts=12,fgm=6)])
        with patch('nbabot.master_source.validate',return_value={'status':'failed','checks':[{'severity':'error','code':'fixture'}]}):
            with self.assertRaisesRegex(ValueError,'Master validation failed'):import_master(self.store,new)
        self.assertEqual(master_status(self.store)['file_hash'],before['file_hash'])
        self.assertEqual(self.store.one('SELECT pts FROM player_games')['pts'],10)
        self.assertTrue(master_path(self.store).is_file())

    def test_new_master_replaces_old_source_and_old_feature_versions(self):
        import_master(self.store,self.organized())
        self.store.snapshot('rosters','sentinel','2025-01-01T00:00:00Z','uploaded-old',{'team':'OTHER'})
        self.store.db.commit()
        new=self.organized([self.row(game_id='0022400002',date_et='2025-01-02')])
        import_master(self.store,new)
        self.assertEqual(self.store.rows('SELECT game_id FROM player_games'),[{'game_id':'0022400002'}])
        self.assertEqual(self.store.rows('SELECT game_id FROM stat_versions'),[{'game_id':'0022400002'}])
        self.assertEqual(self.store.one('SELECT count(*) n FROM snapshots')['n'],0)
        self.assertEqual(len(list((self.store.root/'originals').glob('*.xlsx'))),2)

    def test_schedule_table_does_not_ingest_role_table_as_games(self):
        context={'game_id':'0022699999','date_et':'2026-10-30','tipoff':'2026-10-30T23:00:00Z','home':'NYK','away':'BOS','status':'postponed','season':'2026-27','known_at':'2026-10-01T00:00:00Z'}
        import_master(self.store,self.organized(context=context))
        self.assertEqual(self.store.one("SELECT status FROM games WHERE id='0022699999'")['status'],'postponed')
        self.assertIsNone(self.store.one("SELECT * FROM games WHERE id='wrong-role-game'"))
        self.assertEqual(self.store.one('SELECT count(*) n FROM freshness')['n'],0)

    def test_dnp_not_a_played_zero_and_exact_text_conflict_rejected(self):
        path=self.organized([self.row(appeared=0,minutes=0,min_seconds=0,minutes_exact_text='0',min_seconds_exact_text='0',pts=0,fgm=0)])
        import_master(self.store,path)
        row=self.store.one('SELECT appeared,minutes,pts FROM player_games')
        self.assertEqual(row,{'appeared':0,'minutes':None,'pts':None})
        bad=self.organized([self.row(minutes_exact_text='35')])
        with self.assertRaisesRegex(ValueError,'Exact and numeric'):import_master(self.store,bad)

    def test_refresh_recalculates_without_external_fetch_or_freshness_change(self):
        self.assertEqual(refresh_master(self.store)['status'],'waiting_for_master')
        import_master(self.store,self.organized())
        before=master_status(self.store)['uploaded_at']
        with (patch('nbabot.maintenance.fetch_json',side_effect=AssertionError('Network forbidden')),
             patch('nbabot.model.project_upcoming',return_value=[]) as project,
             patch('nbabot.report.write_report',return_value={'status':'written'})):
            first=refresh_master(self.store);second=refresh_master(self.store)
        self.assertEqual(project.call_count,2)
        self.assertTrue(first['recalc_done']);self.assertEqual(second['status'],'unchanged')
        self.assertEqual(master_status(self.store)['uploaded_at'],before)
        self.assertEqual(self.store.one('SELECT count(*) n FROM freshness')['n'],0)

    def test_mode_required_and_tampered_file_refused(self):
        self.config['data_source_mode']='external'
        with self.assertRaisesRegex(ValueError,'spreadsheet'):import_master(self.store,self.organized())
        self.config['data_source_mode']='spreadsheet'
        import_master(self.store,self.organized())
        master_path(self.store).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError,'missing or has changed'):refresh_master(self.store)

    def test_original_schedule_alias_deduplicates_completed_fixture(self):
        path=self.organized()
        from openpyxl import load_workbook
        book=load_workbook(path)
        sheet=book.create_sheet('Schedule')
        sheet.append(['Date (ET)','Tipoff (ET)','ESPN game ID','Home / away','Opponent','Status'])
        sheet.append(['2025-01-01','7:00 PM','401000001','Home','BOS','Completed'])
        book.save(path);book.close()
        import_master(self.store,path)
        self.assertEqual(self.store.one('SELECT count(*) n FROM games')['n'],1)
        self.assertEqual(self.store.one("SELECT game_id FROM game_aliases WHERE alias='espn:401000001'")['game_id'],'0022400001')

    def test_wrong_season_or_unsupported_id_rejected(self):
        for row in [self.row(season='2025-26'),self.row(player_id='invented-name-id')]:
            with self.assertRaises(ValueError):import_master(self.store,self.organized([row]))

    def test_flat_context_stops_at_separator_before_observed_roles(self):
        path=self.organized()
        from openpyxl import load_workbook
        book=load_workbook(path);sheet=book.create_sheet('Data_GameContext')
        sheet.append(['game_id','date_et','status','home','away'])
        sheet.append(['0022699999','2026-10-30','scheduled','NYK','BOS'])
        sheet.append([]);sheet.append(['game_id','player_id','date_et','started'])
        sheet.append(['0042400011','00456','2025-05-01',1])
        book.save(path);book.close()
        import_master(self.store,path)
        self.assertIsNotNone(self.store.one("SELECT id FROM games WHERE id='0022699999'"))
        self.assertIsNone(self.store.one("SELECT id FROM games WHERE id='0042400011'"))

    def test_unsafe_and_oversized_archives_rejected_before_active_changes(self):
        import_master(self.store,self.organized())
        previous=master_status(self.store)['file_hash']
        bad=self.root/'unsafe.xlsx'
        with zipfile.ZipFile(bad,'w') as z:z.writestr('../escape','bad')
        with self.assertRaisesRegex(ValueError,'Unsafe'):import_master(self.store,bad)
        with patch('nbabot.master_source.MAX_FILE_BYTES',1):
            with self.assertRaisesRegex(ValueError,'32MiB'):import_master(self.store,self.organized())
        self.assertEqual(master_status(self.store)['file_hash'],previous)

    def test_explicit_supplied_team_ids_do_not_relabel_historical_players(self):
        path=self.organized()
        from openpyxl import load_workbook
        book=load_workbook(path);sheet=book.create_sheet('Data_TeamGames')
        sheet.append(['TEAM_ID','TEAM_ABBREVIATION','GAME_ID'])
        sheet.append(['1610612752','NYK','0022400001'])
        sheet.append(['1610612738','BOS','0022400001'])
        book.save(path);book.close()
        import_master(self.store,path)
        self.assertEqual(self.store.one("SELECT abbr FROM teams WHERE id='1610612752'")['abbr'],'NYK')
        self.assertEqual(self.store.one('SELECT team FROM player_games')['team'],'NYK')


if __name__=='__main__':unittest.main()

import io
import json
import os
import tempfile
import unittest
from datetime import datetime,timedelta,timezone
from pathlib import Path
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook

from nbabot.core import Store,load_config
from nbabot.maintenance import fetch_json,refresh,source_check
from nbabot.web import create_app


def master_bytes(game_id="0022600001",points=21):
    workbook=Workbook();sheet=workbook.active;sheet.title="Clean Stats"
    sheet.append(['GAME_ID','PLAYER_ID','PLAYER_NAME','TEAM_ABBREVIATION','Opponent','Season','Game date','Appeared','MIN','PTS','REB','AST','FG3M','FG3A','FGM','FGA','FTM','FTA'])
    played=(datetime.now(timezone.utc)-timedelta(days=2)).date().isoformat()
    sheet.append([game_id,'00123','Test Guard','NYK','BOS','2026-27',played,'Yes',34.098333333333336,points,0,None,1,3,6,19,8,8])
    stream=io.BytesIO();workbook.save(stream);return stream.getvalue()


class SpreadsheetWeb(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.config=load_config(Path(self.tmp.name)/'none')
        self.config.update(data_dir=str(Path(self.tmp.name)/'data'),data_source_mode='spreadsheet',organized_players=[],simulations=100)
        self.env=patch.dict(os.environ,{'NBABOT_DASHBOARD_USER':'owner','NBABOT_DASHBOARD_PASSWORD':'test-only-fixture-password'})
        self.env.start()
        self.auth=('owner','test-only-fixture-password')

    def tearDown(self):
        self.env.stop();self.tmp.cleanup()

    def test_upload_is_private_and_preserves_the_exact_master(self):
        body=master_bytes()
        supplied=load_workbook(io.BytesIO(body),read_only=True,data_only=True)
        source_minutes=supplied.active['I2'].value
        supplied.close()
        with patch('nbabot.maintenance.httpx.Client',side_effect=AssertionError('External data fetch forbidden')):
            with TestClient(create_app(self.config)) as client:
                self.assertEqual(client.post('/api/master',content=body).status_code,401)
                self.assertEqual(client.get('/api/master/download').status_code,401)
                self.assertEqual(client.post('/api/master',content=body,auth=self.auth).status_code,403)
                response=client.post('/api/master',content=body,auth=self.auth,headers={'X-NBAPPBOT-Upload':'master','X-File-Name':'Master.xlsx'})
                self.assertEqual(response.status_code,202,response.text)
                overview=client.get('/api/overview',auth=self.auth).json()
                self.assertEqual(overview['data_source_mode'],'spreadsheet')
                self.assertEqual(overview['initialization']['status'],'ready',overview['initialization'])
                self.assertEqual(overview['master_source']['status'],'ready')
                self.assertEqual(client.get('/api/master/download',auth=self.auth).content,body)
                self.assertEqual(client.get('/api/export/Knicks_Historical_Organized.xlsx',auth=self.auth).content,body)
                with StoreContext(self.config) as store:
                    row=store.one('SELECT game_id,player_id,minutes,reb,ast FROM player_games')
                    self.assertEqual(row['game_id'],'0022600001')
                    self.assertEqual(row['player_id'],'00123')
                    self.assertEqual(row['minutes'],source_minutes)
                    self.assertEqual(row['reb'],0);self.assertIsNone(row['ast'])
                    self.assertEqual(store.rows('SELECT * FROM freshness'),[])

    def test_invalid_upload_retains_last_valid_master(self):
        body=master_bytes()
        headers={'X-NBAPPBOT-Upload':'master','X-File-Name':'Master.xlsx'}
        with TestClient(create_app(self.config)) as client:
            client.post('/api/master',content=body,auth=self.auth,headers=headers)
            client.post('/api/master',content=master_bytes(points=999),auth=self.auth,headers=headers)
            info=client.get('/api/overview',auth=self.auth).json()
            self.assertEqual(info['initialization']['status'],'error')
            self.assertEqual(info['master_source']['status'],'ready')
            self.assertEqual(client.get('/api/master/download',auth=self.auth).content,body)
            self.assertEqual(client.post('/api/master',content=b'bad',auth=self.auth,headers={**headers,'X-File-Name':'file.csv'}).status_code,415)
            self.assertEqual(client.post('/api/master',content=b'',auth=self.auth,headers=headers).status_code,400)

    def test_network_is_hard_disabled_even_if_old_flags_are_enabled(self):
        with StoreContext(self.config) as store,patch('httpx.Client.get',side_effect=AssertionError('Network forbidden')):
            store.config.update(nba_enabled=True,public_context_enabled=True,provider_url='https://example.test/feed')
            with self.assertRaisesRegex(ValueError,'approval'):
                fetch_json('https://stats.nba.com/stats/playergamelogs',store.config)
            self.assertEqual(source_check(store)['status'],'awaiting_approval')
            self.assertEqual(refresh(store)['status'],'waiting_for_master')

    def test_spreadsheet_mode_cannot_start_without_auth(self):
        with patch.dict(os.environ,{'NBABOT_DASHBOARD_USER':'','NBABOT_DASHBOARD_PASSWORD':''}):
            with self.assertRaises(ValueError):create_app(self.config)

    def test_interrupted_upload_can_be_retried_after_restart(self):
        from nbabot.first_boot import write_status
        write_status(self.config['data_dir'],'validating_master',job_id='interrupted')
        with TestClient(create_app(self.config)) as client:
            info=client.get('/api/overview',auth=self.auth).json()
            self.assertEqual(info['initialization']['error_type'],'InterruptedProcessing')
            response=client.post('/api/master',content=master_bytes(),auth=self.auth,
                headers={'X-NBAPPBOT-Upload':'master','X-File-Name':'Master.xlsx'})
            self.assertEqual(response.status_code,202)
            self.assertEqual(client.get('/api/overview',auth=self.auth).json()['initialization']['status'],'ready')

    def test_approved_source_is_available_but_approval_does_not_cover_other_hosts(self):
        self.config.update(external_sources_approved=True,external_source_approvals=['stats.nba.com'])
        response=unittest.mock.MagicMock()
        response.content=b'{}';response.json.return_value={'ok':True}
        response.raise_for_status.return_value=None
        client=unittest.mock.MagicMock()
        client.__enter__.return_value.get.return_value=response
        with patch('nbabot.maintenance.httpx.Client',return_value=client) as factory:
            self.assertEqual(fetch_json('https://stats.nba.com/stats/playergamelogs',self.config),{'ok':True})
            self.assertEqual(factory.call_count,1)
            with self.assertRaisesRegex(ValueError,'approval'):
                fetch_json('https://site.api.espn.com/apis/site/v2/sports/basketball/nba/news',self.config)
            self.assertEqual(factory.call_count,1)


class StoreContext:
    def __init__(self,config):self.config=config
    def __enter__(self):self.store=Store(self.config);return self.store
    def __exit__(self,*args):self.store.close()


if __name__=='__main__':unittest.main()

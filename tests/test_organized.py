import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.utils.cell import range_boundaries

from nbabot.core import Store, dumps, load_config
from nbabot.ingest import normalize_game, upsert_player_game
from nbabot.organized import (SUPPORTING_SHEETS, _summary, _scoped_correction_audit, archive_audit,
                              build_organized_workbook, resolve_selected_players)


class OrganizedWorkbook(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.config = load_config(self.root / "missing.json")
        self.config.update(data_dir=str(self.root / "data"), nba_enabled=False,
                           public_context_enabled=False, season="2025-26",
                           organized_players=["00" + str(i) for i in range(1, 8)])
        self.store = Store(self.config)
        self.at = "2025-02-01T00:00:00+00:00"
        for i in range(1, 9):
            for day in (1, 2):
                team = "NYK" if i <= 7 else "BOS"
                if i == 1 and day == 2:
                    team = "BOS"  # A later trade does not relabel game one.
                self.insert(i, day, team=team)
        self.insert(1, 3, team="NYK", Appeared="No", MIN=0, PTS=0, FGM=0, FGA=0, FG3A=0)
        self.store.db.commit()

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def insert(self, player, day, team="NYK", **extra):
        row = {"GAME_ID": "002250000" + str(day), "PLAYER_ID": "00" + str(player),
               "PLAYER_NAME": "Test Player" + str(player), "TEAM_ABBREVIATION": team,
               "Opponent": "BOS" if team == "NYK" else "NYK", "Season": "2025-26",
               "Game date": "2025-01-0" + str(day), "Appeared": "Yes", "MIN": 34.098333333333336,
               "PTS": 10, "REB": 0, "AST": None if day == 1 else 0, "OREB": 0, "DREB": 0,
               "FGM": 5, "FGA": 11, "FG3M": 0, "FG3A": 1, "FTM": 0, "FTA": 0,
               "Source ID": "https://example.test/official"}
        row.update(extra)
        upsert_player_game(self.store, normalize_game(row, self.at))

    def minimal_audit(self):
        workbook = Workbook()
        source = workbook.active
        source.title = "Source_catalog"
        source.append(["dataset", "source_url", "retrieved_at"])
        source.append(["passing_single_date_check", "https://example.test/tracking", self.at])
        coverage = workbook.create_sheet("Opponent_coverage")
        coverage.append(["team", "named_players"])
        coverage.append(["BOS", 1])
        tracking = workbook.create_sheet("Passing_one_date_2025")
        tracking.append(["PLAYER_ID", "TEAM_ID", "TEAM_ABBREVIATION", "ACTUAL_HISTORICAL_TEAM", "SOURCE_TEAM_LABEL_DIFFERS", "DERIVED_GAME_ID"])
        tracking.append(["001", "1610612738", "BOS", "NYK", True, "0022500001"])
        path = self.root / "Historical_Data_Audit.xlsx"
        workbook.save(path)
        return path

    def test_seven_selection_is_explicit_unique_and_stable(self):
        self.assertEqual([p["id"] for p in resolve_selected_players(self.store)], self.config["organized_players"])
        self.assertEqual(resolve_selected_players(self.store, ["Test Player" + str(i) for i in range(1, 8)])[0]["id"], "001")
        with self.assertRaisesRegex(ValueError, "exactly seven"):
            resolve_selected_players(self.store, [])
        with self.assertRaisesRegex(ValueError, "unique"):
            resolve_selected_players(self.store, ["001"] * 7)
        with self.assertRaisesRegex(ValueError, "Knicks history"):
            resolve_selected_players(self.store, ["00" + str(i) for i in range(2, 9)])

    def test_summaries_keep_missing_zero_and_dnp_separate(self):
        games = self.store.rows("SELECT * FROM player_games WHERE player_id='001'")
        result = _summary(games)
        self.assertEqual(result["played_games"], 2)
        self.assertEqual(result["listed_nonappearances"], 1)
        self.assertEqual(result["reb_total"], 0)
        self.assertEqual(result["ast_observed"], 1)
        self.assertEqual(result["ast_mean_observed"], 0)
        self.assertIsNone(result["ast_total"])
        self.assertEqual(result["actual_teams"], "BOS, NYK")

    def test_eight_player_cohort_keeps_drummond_and_dynamic_sheet_order(self):
        self.insert(8, 3, team="NYK")
        self.store.db.execute("UPDATE players SET name='Andre Drummond' WHERE id='008'")
        self.store.db.commit()
        selection = ["00" + str(i) for i in range(1, 9)]
        report = build_organized_workbook(self.store, self.root / "eight", selected_players=selection)
        self.assertEqual(len(report["selected_players"]), 8)
        self.assertEqual(len(report["sheets"]), 53)
        self.assertEqual(report["sheets"][7], "NYK_Drummond")
        self.assertTrue(all(name.startswith("NYK_vs_") for name in report["sheets"][8:37]))
        self.assertEqual(report["sheets"][37:], SUPPORTING_SHEETS)
        changelog = (self.root / "eight" / "organized_changelog.md").read_text()
        self.assertIn("Created 8 explicitly selected player pages", changelog)
        self.assertEqual(self.config["organized_players"], selection[:7])  # A one-off override never changes saved selection.

    def test_workbook_layout_tables_ids_minutes_and_trade(self):
        audit = self.minimal_audit()
        original_bytes = audit.read_bytes()
        archive_audit(self.store, audit)
        output = self.root / "output"
        report = build_organized_workbook(self.store, output, csv_dir=output / "csv")
        self.assertEqual(audit.read_bytes(), original_bytes)
        self.assertEqual(report["canonical_unique_player_games"], 17)
        self.assertEqual(len(report["sheets"]), 52)
        self.assertEqual(report["sheets"][36:], SUPPORTING_SHEETS)
        workbook = load_workbook(output / "Knicks_Historical_Organized.xlsx", data_only=False)
        try:
            self.assertEqual(workbook.sheetnames[:7], ["NYK_Player" + str(i) for i in range(1, 8)])
            self.assertTrue(all(name.startswith("NYK_vs_") for name in workbook.sheetnames[7:36]))
            for sheet in workbook:
                self.assertEqual(sheet.freeze_panes, "D7")
                ranges = sorted((range_boundaries(t.ref) for t in sheet.tables.values()), key=lambda r: r[1])
                self.assertTrue(all(previous[3] < following[1] for previous, following in zip(ranges, ranges[1:])))
                self.assertTrue(all(t.autoFilter is not None for t in sheet.tables.values()))
            sheet = workbook["Data_PlayerGames"]
            table = next(iter(sheet.tables.values()))
            start = range_boundaries(table.ref)[1]
            headers = [cell.value for cell in sheet[start]]
            row = list(sheet[start + 1])
            self.assertEqual(row[headers.index("game_id")].value, "0022500001")
            self.assertEqual(row[headers.index("game_id")].data_type, "s")
            self.assertEqual(row[headers.index("player_id")].value, "001")
            self.assertAlmostEqual(row[headers.index("minutes")].value, 34.098333333333336)
            self.assertEqual(row[headers.index("minutes_exact_text")].value, "34.098333333333336")
            self.assertEqual(row[headers.index("min_seconds")].value, 2045.9)
            self.assertEqual(row[headers.index("reb")].value, 0)
            self.assertIsNone(row[headers.index("ast")].value)
            records = [dict(zip(headers, r)) for r in sheet.iter_rows(min_row=start + 1, values_only=True)]
            player = [r for r in records if r["player_id"] == "001"]
            self.assertEqual([r["team"] for r in player[:2]], ["NYK", "BOS"])
            tracking = workbook["Data_TrackingGames"]
            raw_values = list(tracking.values)
            header_row = next(i for i, r in enumerate(raw_values) if "SOURCE_TEAM_LABEL_DIFFERS" in r)
            heads = list(raw_values[header_row])
            values = raw_values[header_row + 1]
            self.assertEqual(values[heads.index("TEAM_ABBREVIATION")], "BOS")
            self.assertEqual(values[heads.index("ACTUAL_HISTORICAL_TEAM")], "NYK")
            self.assertEqual(values[heads.index("source_url")], "https://example.test/tracking")
            self.assertIsNotNone(workbook["Contents"]["A6"].hyperlink)
        finally:
            workbook.close()
        self.assertTrue((output / "organized_changelog.md").is_file())
        self.assertTrue((output / "organized_missing.csv").is_file())
        self.assertEqual(json.loads((output / "organized_validation.json").read_text())["canonical_player_games"], 17)

    def test_corrupt_reference_and_failed_validation_cannot_replace_output(self):
        audit = self.minimal_audit()
        manifest = archive_audit(self.store, audit)
        target = self.store.root / "reference" / manifest["file"]
        target.write_bytes(b"corrupt")
        output = self.root / "Knicks_Historical_Organized.xlsx"
        output.write_bytes(b"previous valid output")
        with self.assertRaisesRegex(ValueError, "checksum"):
            build_organized_workbook(self.store, output)
        self.assertEqual(output.read_bytes(), b"previous valid output")
        with patch("nbabot.ingest.validate", return_value={"status": "failed", "checks": []}):
            with self.assertRaisesRegex(ValueError, "validation"):
                build_organized_workbook(self.store, output)
        self.assertEqual(output.read_bytes(), b"previous valid output")

    def test_updated_context_stays_in_separate_units_and_scopes(self):
        context = {"tables": [{"key": "2025-26:Passing", "kind": "tracking_season", "measure": "Passing", "season": "2025-26",
                    "scope": "Season per-game means", "unit": "PerGame", "source_url": "https://example.test/current",
                    "retrieved_at": self.at, "coverage_cutoff": None, "refresh_status": "success",
                    "rows": [{"PLAYER_ID": "001", "PLAYER_NAME": "Test Player1", "POTENTIAL_AST": 8.5, "GP": 2}]}], "errors": []}
        with patch("nbabot.historical_context.load_context_reference", return_value=context):
            report = build_organized_workbook(self.store, self.root / "new")
        self.assertEqual(report["canonical_player_games"], 17)
        tracking_table = [t for t in report["tables"] if t["sheet"] == "Data_TrackingSeason" and t["table"] == "Season tracking 2025-26 Passing"]
        self.assertEqual(tracking_table[0]["rows"], 1)

    def test_fresh_public_data_has_no_private_audit_counts_or_assumed_rules(self):
        output=self.root/'public-first-start'
        report=build_organized_workbook(self.store,output)
        warnings={row['code']:row for row in report['unavailable_and_cautions']}
        self.assertEqual(warnings['audit_reference']['availability'],'Unavailable')
        self.assertEqual(warnings['named_matchup_scope']['availability'],'Unavailable')
        self.assertIn('private uploaded workbook',warnings['named_matchup_scope']['detail'])
        self.assertIn('not a verified zero-gap result',warnings['unavailable_defender_pairs']['detail'])
        self.assertNotIn('lineup_cap',warnings)
        self.assertNotIn('tracking_team_labels',warnings)
        self.assertNotIn('lineup_minute_gaps',warnings)
        text=json.dumps(report['unavailable_and_cautions'])
        for stale_claim in ('77 regular-season','42 one-date','19,514','all19514'):
            self.assertNotIn(stale_claim,text)
        book=load_workbook(output/'Knicks_Historical_Organized.xlsx',read_only=True,data_only=True)
        try:
            values=' '.join(str(value) for row in book['Data_LineupsSeason'].values for value in row if value is not None)
            self.assertIn('No archived season lineup combinations were supplied',values)
            self.assertNotIn('19514',values)
        finally:
            book.close()

    def test_named_pairs_exclude_unrelated_roster_players_and_other_opponent_games(self):
        self.insert(9, 1, team="BOS")
        self.store.db.execute("INSERT INTO raw_rows VALUES(?,?,?,?)", ("scope-fixture", "Matchups", 1,
            dumps({"Knicks player":"Test Player1", "Opponent":"BOS", "Suspected defender":"Test Player8",
                   "Alternate defender":None, "Assignment status":"Estimated"})))
        # Player8's second canonical game is a broad baseline against another
        # opponent. It must not become a Knicks head-to-head observation.
        self.store.db.execute("UPDATE player_games SET opponent='ORL' WHERE player_id='008' AND date_et='2025-01-02'")
        self.store.db.commit()
        output=self.root/"scoped"
        report=build_organized_workbook(self.store,output)
        self.assertEqual({p['opponent_id'] for p in report['scope']['pairs']},{'008'})
        workbook=load_workbook(output/'Knicks_Historical_Organized.xlsx',data_only=True)
        try:
            sheet=workbook['Data_PlayerGames']
            table=next(iter(sheet.tables.values()))
            start=range_boundaries(table.ref)[1]
            headers=[c.value for c in sheet[start]]
            data=[dict(zip(headers,r)) for r in sheet.iter_rows(min_row=start+1,max_row=range_boundaries(table.ref)[3],values_only=True)]
            self.assertNotIn('009',{r['player_id'] for r in data})
            opponent=[r for r in data if r['player_id']=='008']
            self.assertEqual(len(opponent),1)
            self.assertEqual(opponent[0]['opponent'],'NYK')
            pair_table=next(t for t in report['tables'] if t['sheet']=='NYK_Player1' and t['table'].startswith('Specific named matchup pairs'))
            self.assertEqual(pair_table['rows'],1)
            pair_sheet=workbook['NYK_Player1']
            pair_start=range_boundaries(pair_table['range'])[1]
            pair_headers=[c.value for c in pair_sheet[pair_start]]
            pair=dict(zip(pair_headers,[c.value for c in pair_sheet[pair_start+1]]))
            self.assertEqual(pair['opponent_full_game_samples'],1)
            self.assertEqual(pair['opponent_baseline_played_games'],2)
        finally:
            workbook.close()
        self.assertEqual(self.store.one('SELECT count(*) n FROM player_games')['n'],18)

    def test_visible_audit_is_limited_to_exact_exported_game_player_keys(self):
        # The same named opponent can have many other fixtures and annual
        # corrections; those payloads do not belong in this matchup workbook.
        rows=[{'game_id':'0022500001','player_id':'008'},
              {'game_id':'0042400001','player_id':'001'}]
        records=[{'record_key':key,'old':'retained original','new':'published correction'} for key in (
            '0022500001/008', '0022500002/008', '2025-26/008',
            '0022500001/009', '0022500001', '0022599999',
            '0042400001/001', '0022500001/008/extra')]
        selected=_scoped_correction_audit(records,rows)
        self.assertEqual([r['record_key'] for r in selected],
                         ['0022500001/008','0022500001','0042400001/001'])
        self.assertEqual(len(records),8)  # Filtering never deletes the audit archive.

    def test_career_phase_dedup_and_new_role_knowledge_stay_separate(self):
        canonical=self.store.one("SELECT * FROM player_games WHERE player_id='001' AND date_et='2025-01-01'")
        old={**canonical,'game_id':'0022300001','season':'2023-24','date_et':'2023-10-30',
             'season_type':'Regular Season','pts':14,'known_at':'2026-10-26T20:00:00+00:00','source':'https://example.test/career'}
        playoff={**old,'game_id':'0042300001','date_et':'2024-04-20','season_type':'Playoffs','pts':30}
        overlap={**canonical,'season_type':'Regular Season','minutes':41.9,'source_minute_clock':'41:54',
                 'known_at':'2026-10-26T20:00:00+00:00','source':'https://example.test/newer-career'}
        career={'rows':[old,playoff,overlap],'tables':[],'controls':[],'errors':[],'status':'success'}
        starts={'tracking_games':[],'passing_games':[],'starts':[{'game_id':canonical['game_id'],'player_id':'001',
            'started':1,'validation_status':'source_conflict',
            'role_validation_status':'verified_official_game_identity_and_five_starters_per_team',
            'flags':['source_minute_difference'],'source_url':'https://example.test/box',
            'retrieved_at':'2026-10-26T21:00:00+00:00'}],'sources':[],'checks':[],'errors':[],'coverage':{}}
        with patch('nbabot.career_reference.load_career_reference',return_value=career), \
             patch('nbabot.opportunity_reference.load_opportunity_reference',return_value=starts):
            report=build_organized_workbook(self.store,self.root/'career')
        workbook=load_workbook(self.root/'career/Knicks_Historical_Organized.xlsx',data_only=True)
        try:
            sheet=workbook['Data_PlayerGames']; table=next(iter(sheet.tables.values()))
            start=range_boundaries(table.ref)[1]; headers=[c.value for c in sheet[start]]
            rows=[dict(zip(headers,r)) for r in sheet.iter_rows(min_row=start+1,max_row=range_boundaries(table.ref)[3],values_only=True)]
            selected=[r for r in rows if r['player_id']=='001']
            self.assertEqual(len({(r['game_id'],r['player_id']) for r in selected}),len(selected))
            same=next(r for r in selected if r['game_id']==canonical['game_id'])
            self.assertEqual(same['minutes_exact_text'],'34.098333333333336')
            self.assertEqual(same['started'],1)
            self.assertEqual(same['known_at'],canonical['known_at'])
            self.assertEqual(same['started_known_at'],'2026-10-26T21:00:00+00:00')
            self.assertIn('source_minute_difference',same['started_source_flags'])
            season=workbook['Data_PlayerSeason']; first=next(iter(season.tables.values()))
            begin=range_boundaries(first.ref)[1]; heads=[c.value for c in season[begin]]
            annual=[dict(zip(heads,r)) for r in season.iter_rows(min_row=begin+1,max_row=range_boundaries(first.ref)[3],values_only=True)]
            older=[r for r in annual if r['player_id']=='001' and r['season']=='2023-24']
            self.assertEqual({r['season_type']:r['pts_total'] for r in older},{'Regular Season':14,'Playoffs':30})
        finally:
            workbook.close()
        self.assertEqual(self.store.one("SELECT started FROM player_games WHERE player_id='001' AND date_et='2025-01-01'")['started'],None)
        self.assertEqual(report['canonical_player_games'],17)

    def test_named_opponent_playoff_box_is_added_without_mutating_canonical(self):
        self.store.db.execute("INSERT INTO raw_rows VALUES(?,?,?,?)", ('playoff-scope','Matchups',1,
            dumps({'Knicks player':'Test Player1','Opponent':'BOS','Suspected defender':'Test Player8'})))
        self.store.db.commit()
        row=self.store.one("SELECT pg.*,p.name FROM player_games pg JOIN players p ON p.id=pg.player_id WHERE player_id='008' LIMIT 1")
        row.update(game_id='0042400001',date_et='2025-05-01',season='2024-25',season_type='Playoffs',
            known_at='2026-10-26T21:00:00+00:00',retrieved_at='2026-10-26T21:00:00+00:00',
            source_url='https://example.test/playoff-box',validation_status='source_game_identity_only',
            started=1,role_validation_status='verified_official_game_identity_and_five_starters_per_team')
        source={'tracking_games':[],'passing_games':[],'starts':[row],'sources':[],'checks':[],'errors':[],'coverage':{}}
        with patch('nbabot.opportunity_reference.load_opportunity_reference',return_value=source):
            report=build_organized_workbook(self.store,self.root/'playoff-box')
        book=load_workbook(self.root/'playoff-box/Knicks_Historical_Organized.xlsx',data_only=True)
        try:
            sheet=book['Data_PlayerGames']; table=next(iter(sheet.tables.values()))
            start=range_boundaries(table.ref)[1]; headers=[c.value for c in sheet[start]]
            rows=[dict(zip(headers,r)) for r in sheet.iter_rows(min_row=start+1,max_row=range_boundaries(table.ref)[3],values_only=True)]
            added=next(r for r in rows if r['game_id']=='0042400001')
            self.assertEqual(added['player_id'],'008')
            self.assertEqual(added['season_type'],'Playoffs')
            self.assertEqual(added['pts'],10)
            self.assertEqual(added['source'],'https://example.test/playoff-box')
        finally:
            book.close()
        self.assertEqual(report['canonical_player_games'],17)
        self.assertIsNone(self.store.one("SELECT game_id FROM player_games WHERE game_id='0042400001'"))


if __name__ == "__main__":
    unittest.main()

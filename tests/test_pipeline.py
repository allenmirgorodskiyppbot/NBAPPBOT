import csv
import json
import tempfile
import unittest
from datetime import datetime, timedelta, timezone
from pathlib import Path
from unittest.mock import patch

import numpy as np
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from nbabot.core import Store, digest, dumps, freshness, load_config, now, stamp
from nbabot.evaluation import evaluate
from nbabot.ingest import import_files, normalize_game, upsert_game, upsert_player_game, validate
from nbabot.maintenance import export_package, refresh
from nbabot.model import history, probabilities, project_upcoming, roster_for, simulate
from nbabot.offers import (bankroll, build_entries, current_offers, expected_value, import_board,
                           import_rules, payout_multiplier, rank, settle_entry)
from nbabot.scheduler import plan_jobs, tick
from nbabot.web import create_app


class Pipeline(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.path=Path(self.tmp.name)
        self.config=load_config(self.path/"none")
        self.config.update(data_dir=str(self.path/"data"),nba_enabled=False,public_context_enabled=False,simulations=400)
        self.store=Store(self.config)
        self.old=stamp(datetime.now(timezone.utc)-timedelta(days=90))
        self.tip=stamp(datetime.now(timezone.utc)+timedelta(days=1))

    def tearDown(self):
        self.store.close(); self.tmp.cleanup()

    def row(self,**kw):
        r={"GAME_ID":"0022600001","PLAYER_ID":"00123","PLAYER_NAME":"Test Guard",
           "TEAM_ABBREVIATION":"NYK","Opponent":"BOS","Season":"2026-27","Game date":"2026-01-01",
           "Appeared":"Yes","MIN":34.098333333333336,"PTS":21,"REB":5,"AST":3,
           "FG3M":1,"FG3A":3,"FGM":6,"FGA":19,"FTM":8,"FTA":8,"OREB":0,"DREB":5}
        r.update(kw); return r

    def insert(self,row=None,known=None):
        p=normalize_game(row or self.row(),known or self.old)
        result=upsert_player_game(self.store,p)
        self.store.db.commit(); return result

    def schedule(self,status="scheduled"):
        upsert_game(self.store,{"id":"0022600099","date_et":self.tip[:10],"tipoff":self.tip,"home":"BOS","away":"NYK",
                               "status":status,"known_at":now(),"source":"https://example.test/schedule"})
        self.store.db.commit()

    def many(self,n=45):
        start=datetime.now(timezone.utc)-timedelta(days=n+4)
        for i in range(n):
            for pid,team,opp in [("00123","NYK","BOS"),("00456","BOS","NYK")]:
                r=self.row(GAME_ID=f"002260{i:04}",PLAYER_ID=pid,PLAYER_NAME=pid,TEAM_ABBREVIATION=team,Opponent=opp,
                           **{"Game date":(start+timedelta(days=i)).date().isoformat()})
                r["PTS"]=20+i%5; r["FGM"]=None; r["FTM"]=None
                self.insert(r)
        self.schedule()

    def board(self,board_id="one",line=20.5,sides=None,captured=None,players=("00123",)):
        obj={"board_id":board_id,"offers":[{"player_id":pid,"game_id":"0022600099","market":"points","line":line,
            "sides":sides or ["more","less"],"offer_type":"standard","captured_at":captured or now(),
            "source_url":"https://app.prizepicks.com/","contest":"test_only","rules_id":"fixture-rules"} for pid in players]}
        path=self.path/(board_id+".json"); path.write_text(json.dumps(obj)); return path

    def rules(self):
        return {"id":"fixture-rules","verified":True,"verified_at":self.old,"effective_from":self.old,
            "source_url":"https://example.test/test-only-rules","contest":"test_only",
            "eligibility":{"min_legs":2,"max_legs":3,"min_teams":2,"same_game_allowed":True},
            "treatment":{"tie":"remove","dnp":"remove","void":"remove","reboot":"remove"},
            "below_minimum_multiplier":1,"ineligible_reduction_multiplier":1,
            "payouts":{"2":{"2":3,"1":0,"0":0},"3":{"3":5,"2":1,"1":0,"0":0}}}

    def test_precision_missing_zero_and_ids(self):
        p=normalize_game(self.row(AST=None,REB=0,OREB=0,DREB=0),self.old)
        self.assertEqual(p["game_id"],"0022600001"); self.assertEqual(p["player_id"],"00123")
        self.assertIsNone(p["ast"]); self.assertEqual(p["reb"],0)
        self.assertEqual(p["minutes"],34.098333333333336)
        self.assertAlmostEqual(p["min_seconds"],2045.9)

    def test_bad_totals_and_equivalent_fields(self):
        with self.assertRaises(ValueError): normalize_game(self.row(PTS=99),self.old)
        with self.assertRaises(ValueError): normalize_game(self.row(**{"3PM":2}),self.old)
        p=normalize_game(self.row(FG3M=None,**{"3PM":1}),self.old)
        self.assertEqual(p["fg3m"],1)

    def test_incorrect_join_and_source_conflict(self):
        self.insert()
        with self.assertRaises(ValueError): self.insert(self.row(**{"Game date":"2026-01-02","PLAYER_ID":"999"}))
        with self.assertRaises(ValueError): self.insert(self.row(AST=9))

    def test_workbook_csv_dedup_and_raw_preservation(self):
        r=self.row()
        book=Workbook(); sheet=book.active; sheet.title="Clean Stats"
        sheet.append(list(r)); sheet.append(list(r.values()))
        summary=book.create_sheet("Tracking Stats"); summary.append(["Player","GP","PTS"]); summary.append(["Test Guard",40,20])
        x=self.path/"selected.xlsx"; book.save(x)
        c=self.path/"NBA_Player_Games.csv"
        with c.open("w",newline="") as f:
            writer=csv.DictWriter(f,fieldnames=r); writer.writeheader(); writer.writerow(r)
        report=import_files(self.store,[str(x),str(c)])
        self.assertEqual(self.store.one("SELECT count(*) n FROM player_games")["n"],1)
        self.assertEqual(sum(r.get("duplicate",0) for r in report),1)
        self.assertEqual(self.store.one("SELECT count(*) n FROM raw_rows WHERE sheet='Tracking Stats'")["n"],1)
        self.assertTrue((self.path/"data/originals").exists())
        self.assertEqual(import_files(self.store,[str(c)])[0]["status"],"already_imported")

    def test_corrections_keep_asof_version(self):
        self.insert()
        p=normalize_game(self.row(AST=9),now())
        upsert_player_game(self.store,p,corrections=True); self.store.db.commit()
        cutoff=stamp(datetime.now(timezone.utc)-timedelta(days=1))
        self.assertEqual(history(self.store,"00123",cutoff)[0]["ast"],3)
        self.assertEqual(self.store.one("SELECT count(*) n FROM audit WHERE kind='stat_correction'")["n"],1)

    def test_future_data_and_target_minutes_cannot_leak(self):
        self.many()
        cutoff=stamp(datetime.now(timezone.utc)-timedelta(days=20))
        h=history(self.store,"00123",cutoff)
        a=simulate(h,200,seed=5)
        self.insert(self.row(GAME_ID="0099999999",**{"Game date":"2030-01-01","MIN":80}))
        b=simulate(history(self.store,"00123",cutoff),200,seed=5)
        np.testing.assert_array_equal(a["points"],b["points"])
        self.assertTrue(all(r["date_et"]<cutoff[:10] for r in h))

    def test_correction_reversion_preserves_each_knowledge_time(self):
        a=stamp(datetime.now(timezone.utc)-timedelta(days=10))
        b=stamp(datetime.now(timezone.utc)-timedelta(days=5))
        c=stamp(datetime.now(timezone.utc)-timedelta(days=1))
        self.insert(known=a)
        upsert_player_game(self.store,normalize_game(self.row(AST=9),b),corrections=True)
        upsert_player_game(self.store,normalize_game(self.row(AST=3),c),corrections=True); self.store.db.commit()
        cutoff=stamp(datetime.now(timezone.utc)-timedelta(days=3))
        self.assertEqual(history(self.store,"00123",cutoff)[0]["ast"],9)
        self.assertEqual(history(self.store,"00123",now())[0]["ast"],3)

    def test_dnp_is_not_zero_played_game(self):
        self.insert(self.row(Appeared="No",MIN=0,PTS=0,FGM=0,FG3M=0,FTM=0))
        self.assertEqual(history(self.store,"00123",now()),[])

    def test_same_day_final_is_available_only_after_verified_capture(self):
        from nbabot.core import ET
        day=datetime.now(ET).date().isoformat()
        past=stamp(datetime.now(timezone.utc)-timedelta(minutes=1))
        self.insert(self.row(**{"Game date":day,"game_status":"completed"}),known=past)
        self.assertEqual(len(history(self.store,"00123",now())),1)
        before=stamp(datetime.now(timezone.utc)-timedelta(minutes=2))
        self.assertEqual(history(self.store,"00123",before),[])

    def test_lean_source_does_not_erase_known_start(self):
        self.insert(self.row(Started="Yes"))
        new=normalize_game(self.row(AST=9),now())
        upsert_player_game(self.store,new,corrections=True); self.store.db.commit()
        self.assertEqual(self.store.one("SELECT started FROM player_games")["started"],1)

    def test_probability_ties_and_fractional_lines(self):
        self.assertEqual(probabilities([0,1,2,2],2),{"more":0,"less":.5,"tie":.5})
        p=probabilities([0,1,2,2],1.5); self.assertEqual(p["tie"],0); self.assertAlmostEqual(sum(p.values()),1)
        with self.assertRaises(ValueError): probabilities([],2)
        with self.assertRaises(ValueError): probabilities([1],float("nan"))

    def test_threes_attempts_and_shared_minutes(self):
        self.many()
        sims=simulate(history(self.store,"00123",now()),400,seed=7,minute_cap=20)
        self.assertTrue(np.all(sims["minutes"]<=20))
        self.assertTrue(np.all(sims["threes"]<=sims["three_attempts"]))
        self.assertTrue(np.all(sims["points"]>=3*sims["threes"]))

    def test_postponed_game_is_excluded(self):
        self.many(); self.schedule("postponed")
        self.assertEqual(project_upcoming(self.store),[])

    def test_trade_current_roster_not_applied_retrospectively(self):
        self.insert()
        self.store.snapshot("rosters","00123",now(),"https://example.test/roster",{"player_id":"00123","name":"Test Guard","team":"BOS"})
        self.store.db.commit()
        roster,verified=roster_for(self.store,"NYK",now()); self.assertEqual(roster,[])
        self.assertEqual(history(self.store,"00123",now())[0]["team"],"NYK")

    def test_actual_offers_only_and_stale_offer_gate(self):
        self.many(); project_upcoming(self.store)
        self.assertEqual(rank(self.store)["candidates"],[])
        path=self.board(captured=stamp(datetime.now(timezone.utc)-timedelta(hours=4)),sides=["less"])
        import_board(self.store,path); result=rank(self.store)
        self.assertEqual(len(result["candidates"]),1)
        self.assertEqual(result["candidates"][0]["side"],"less")
        self.assertIn("stale_offer",result["candidates"][0]["skip_reasons"])
        self.assertEqual(result["qualifying"],[])

    def test_board_idempotency_and_withdrawn_lines(self):
        self.many()
        p=self.board(); import_board(self.store,p)
        self.assertEqual(import_board(self.store,p)["status"],"already_imported")
        import_board(self.store,self.board("two",line=25.5))
        self.assertEqual([r["line"] for r in current_offers(self.store)],[25.5])

    def test_unreviewed_screenshot_rejected(self):
        self.many(); path=self.board(); obj=json.loads(path.read_text())
        obj["offers"][0]["evidence_type"]="screenshot"; path.write_text(dumps(obj))
        with self.assertRaises(ValueError): import_board(self.store,path)

    def test_no_assumed_independent_joint_entry(self):
        self.many(); project_upcoming(self.store)
        import_board(self.store,self.board(players=("00123","00456")))
        entries=build_entries(self.store,experimental=True)["entries"]
        self.assertTrue(entries)
        for e in entries:
            self.assertEqual({c["team"] for c in e["selections"]},{"NYK","BOS"})
            self.assertIsNone(e["joint_probability"]); self.assertIsNone(e["expected_value"])
            self.assertIn("joint_dependence_not_validated",e["reasons"])

    def test_tie_void_reboot_reduced_settlement(self):
        rules=self.rules()
        self.assertEqual(payout_multiplier(["win","win"],rules),3)
        self.assertEqual(payout_multiplier(["win","tie"],rules),1)
        self.assertEqual(payout_multiplier(["win","win","dnp"],rules,teams=["NYK","BOS","BOS"]),3)
        self.assertEqual(payout_multiplier(["win","win","dnp"],rules,teams=["NYK","NYK","BOS"]),1)
        with self.assertRaises(ValueError): payout_multiplier(["win","win","dnp"],rules)
        self.assertEqual(payout_multiplier(["win","void","reboot"],rules),1)
        self.assertEqual(payout_multiplier(["win","loss"],rules),0)
        result=expected_value([{"probability":.5,"outcomes":["win","win"]},{"probability":.5,"outcomes":["win","loss"]}],rules)
        self.assertEqual(result["expected_net"],.5)
        with self.assertRaises(ValueError): expected_value([{"probability":.5,"outcomes":["win","win"]}],rules)
        with self.assertRaises(ValueError): payout_multiplier(["win","win"],rules,["goblin","standard"])

    def test_paper_bankroll_roi_and_immutable_settlements(self):
        self.many(); project_upcoming(self.store)
        path=self.path/"rules.json"; path.write_text(dumps(self.rules())); import_rules(self.store,path)
        import_board(self.store,self.board(players=("00123","00456")))
        entry=build_entries(self.store,experimental=True)["entries"][0]
        settle_entry(self.store,entry["id"],["win","win"])
        r=bankroll(self.store); self.assertEqual(r["ending_balance"],102); self.assertEqual(r["roi"],2)
        self.assertIsNone(r["roi_95_interval"])
        with self.assertRaises(ValueError): settle_entry(self.store,entry["id"],["loss","loss"])

    def test_package_ids_text_and_missing_values(self):
        self.insert(self.row(AST=None))
        dest=Path(export_package(self.store)); self.assertTrue((dest/"manifest.json").exists())
        w=load_workbook(dest/"Master_Data.xlsx",read_only=True,data_only=True)
        rows=list(w["NBA_Player_Games"].values); h=list(rows[0]); row=rows[1]
        self.assertEqual(row[h.index("game_id")],"0022600001")
        self.assertIsNone(row[h.index("ast")]); self.assertEqual(row[h.index("oreb")],0)
        w.close()
        import sqlite3
        db=sqlite3.connect(dest/"database.sqlite3")
        self.assertEqual(db.execute("SELECT game_id FROM player_games").fetchone()[0],"0022600001"); db.close()

    def test_failed_refresh_retains_verified_files_and_data(self):
        self.insert(); export_package(self.store)
        latest=self.path/"data/packages/latest"; before=(latest/"Master_Data.xlsx").read_bytes()
        result=refresh(self.store)
        self.assertEqual(result["status"],"failed")
        self.assertEqual(before,(latest/"Master_Data.xlsx").read_bytes())
        self.assertEqual(self.store.one("SELECT count(*) n FROM player_games")["n"],1)
        self.assertTrue(freshness(self.store)["stats"]["stale"])

    def test_successful_refresh_repeatable_and_atomic_validation(self):
        self.many(); at=now(); source="https://example.test/authorized-feed"
        stats=[self.row(GAME_ID="0022699998",**{"Game date":(datetime.now(timezone.utc)-timedelta(days=2)).date().isoformat()})]
        schedule=[{"id":"0022600099","date_et":self.tip[:10],"tipoff":self.tip,"home":"BOS","away":"NYK","status":"scheduled"}]
        records={"stats":stats,"schedule":schedule,"rosters":[{"player_id":"00123","name":"Test Guard","team":"NYK","observed_at":at}],
                 "injuries":[],"lineups":[],"news":[]}
        sections={k:{"complete":True,"verified_at":at,"coverage_cutoff":at,"source_url":source,"records":v} for k,v in records.items()}
        feed=self.path/"feed.json"; feed.write_text(dumps({"sections":sections})); self.config["provider_file"]=str(feed)
        first=refresh(self.store); self.assertEqual(first["status"],"success",first)
        count=self.store.one("SELECT count(*) n FROM player_games")["n"]
        second=refresh(self.store); self.assertEqual(second["status"],"success",second)
        self.assertEqual(count,self.store.one("SELECT count(*) n FROM player_games")["n"])
        latest=(self.path/"data/packages/latest").resolve()
        sections["stats"]["records"][0]["PTS"]=999
        feed.write_text(dumps({"sections":sections}))
        third=refresh(self.store); self.assertEqual(third["status"],"failed")
        self.assertEqual(latest,(self.path/"data/packages/latest").resolve())

    def test_daily_schedule_dst_and_retry_idempotency(self):
        jobs=plan_jobs(self.store,"2026-03-07T14:05:00Z")
        self.assertEqual(jobs[0][1],"2026-03-07T14:00:00.000000+00:00")
        jobs=plan_jobs(self.store,"2026-03-08T13:05:00Z")
        self.assertEqual(jobs[0][1],"2026-03-08T13:00:00.000000+00:00")
        at=stamp(datetime.now(timezone.utc).replace(hour=23,minute=0))
        calls=[]
        def runner(s): calls.append(1); return {"status":"success"}
        tick(self.store,at,runner); tick(self.store,at,runner)
        self.assertEqual(len(calls),1)

    def test_heartbeat_survives_concurrent_tick_and_worker_pulse(self):
        from concurrent.futures import ThreadPoolExecutor
        from nbabot.scheduler import heartbeat_file
        with ThreadPoolExecutor(max_workers=4) as pool:
            list(pool.map(lambda _:heartbeat_file(self.store),range(40)))
        recorded=json.loads((self.store.root/"scheduler-heartbeat.json").read_text())
        self.assertTrue(recorded["heartbeat"])

    def test_asof_evaluation_cannot_claim_empty_tests_passed(self):
        self.many()
        # Import all observations today, as in the workbook's extraction vintage.
        self.store.db.execute("UPDATE stat_versions SET known_at=?",(now(),)); self.store.db.commit()
        result=evaluate(self.store,"strict")
        self.assertEqual(result["status"],"unrun_insufficient_asof_history")
        self.assertFalse(result["eligible_for_selection"])
        self.assertTrue(evaluate(self.store,"strict")["cached"])

    def test_new_model_does_not_relabel_viewed_test_dates_as_untouched(self):
        self.many()
        earlier={"mode":"retrospective","status":"completed","created_at":now(),
                 "partitions":{"test_from":"2020-01-01","test_through":"2099-01-01"}}
        self.store.db.execute("INSERT INTO evaluations VALUES('prior',?,?)",(now(),dumps(earlier))); self.store.db.commit()
        result=evaluate(self.store,"retrospective")
        self.assertEqual(result["test_independence"],"previously_viewed_dates")

    def test_dashboard_functional_request_and_exports(self):
        self.insert(); export_package(self.store)
        with TestClient(create_app(self.config)) as client:
            response=client.get("/api/overview"); self.assertEqual(response.status_code,200)
            self.assertEqual(response.json()["status"],"experimental_projections_only")
            self.assertEqual(client.get("/").status_code,200)
            self.assertIn("Projection",client.get("/").text)
            self.assertEqual(client.get("/api/export/NBA_Player_Games.csv").status_code,200)
            self.assertEqual(client.get("/api/export/not-allowed").status_code,404)

    def test_publication_fault_rolls_back_db_and_latest_pointer(self):
        self.insert(); export_package(self.store)
        latest=(self.path/"data/packages/latest").resolve()
        at=now(); source="https://example.test/feed"
        records={"stats":[self.row(GAME_ID="0022699999",**{"Game date":(datetime.now(timezone.utc)-timedelta(days=2)).date().isoformat()})],
            "schedule":[{"id":"0022600099","date_et":self.tip[:10],"tipoff":self.tip,"home":"BOS","away":"NYK","status":"scheduled"}],
            "rosters":[{"player_id":"00123","name":"Test Guard","team":"NYK","observed_at":at}],"injuries":[],"lineups":[],"news":[]}
        obj={"sections":{k:{"complete":True,"verified_at":at,"coverage_cutoff":at,"source_url":source,"records":v} for k,v in records.items()}}
        path=self.path/"feed.json"; path.write_text(dumps(obj)); self.config["provider_file"]=str(path)
        from nbabot.maintenance import publish_package
        def fault(store,dest):
            publish_package(store,dest)
            raise OSError("Simulated disk fault after pointer replacement")
        with patch("nbabot.maintenance.publish_package",fault):
            self.assertEqual(refresh(self.store)["status"],"failed")
        self.assertEqual(latest,(self.path/"data/packages/latest").resolve())
        self.assertEqual(self.store.one("SELECT count(*) n FROM player_games")["n"],1)

    def test_repeated_ranking_supersedes_decisions_but_preserves_archive(self):
        self.many(); project_upcoming(self.store); import_board(self.store,self.board())
        rank(self.store); rank(self.store)
        self.assertEqual(self.store.one("SELECT count(*) n FROM decisions WHERE status='active'")["n"],2)
        self.assertEqual(self.store.one("SELECT count(*) n FROM decisions WHERE status='superseded'")["n"],2)

    def test_unconfirmed_headline_does_not_adjust_projection(self):
        self.many()
        before=project_upcoming(self.store)
        self.store.snapshot("news","rumor",now(),"https://example.test/news",{"headline":"May have a larger role",
            "players":["00123"],"confirmation":"unconfirmed","published_at":now(),"retrieved_at":now()})
        self.store.db.commit(); after=project_upcoming(self.store)
        self.assertEqual(before[0]["markets"],after[0]["markets"])

    def test_scheduler_failed_job_retries_without_duplicate_success(self):
        at=datetime.now(timezone.utc).replace(hour=23,minute=0,second=0,microsecond=0)
        calls=[]
        def runner(s):
            calls.append(1); return {"status":"failed" if len(calls)==1 else "success"}
        tick(self.store,stamp(at),runner)
        tick(self.store,stamp(at+timedelta(seconds=30)),runner)
        self.assertEqual(len(calls),1)
        tick(self.store,stamp(at+timedelta(minutes=2)),runner)
        tick(self.store,stamp(at+timedelta(minutes=3)),runner)
        self.assertEqual(len(calls),2)


if __name__=="__main__":
    unittest.main()

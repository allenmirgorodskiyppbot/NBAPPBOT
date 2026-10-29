"""End-to-end local checks for the persistent app without external credentials."""
import asyncio
from concurrent.futures import ThreadPoolExecutor
from contextlib import suppress
import datetime as dt
from io import BytesIO
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

from fastapi.testclient import TestClient
import openpyxl

import app
from master import MasterStore

PASSWORD = 'a-private-integration-password'
HEADERS = {'X-Chat-Password': PASSWORD}


async def fake_initialize(application):
    application.state.phase = 'ready'
    application.state.detail = 'Test ready'


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)

    def tearDown(self):
        self.directory.cleanup()

    def test_application_selftest_uses_only_temporary_fixtures(self):
        import asyncio
        with patch.object(app, 'DATA_DIR', self.root), patch.object(app, 'CHAT_PASSWORD', PASSWORD), \
                patch.object(app, 'initialize', fake_initialize):
            with TestClient(app.app):
                master = app.app.state.master
                before = master.status()['counts']
                report = asyncio.run(app.self_test())
                self.assertIn('0 failed', report)
                self.assertIn('temporary databases', report)
                self.assertEqual(master.status()['counts'], before)

    def test_web_entry_settlement_download_and_restart(self):
        legs = [{'player': 'Jalen Brunson', 'team': 'NYK', 'stat': 'pts', 'side': 'more', 'line': 20.5,
                 'date': '2026-10-27'}, {'player': 'Derrick White', 'team': 'BOS', 'stat': 'ast',
                 'side': 'less', 'line': 7.5, 'date': '2026-10-27'}]
        with patch.object(app, 'DATA_DIR', self.root), patch.object(app, 'CHAT_PASSWORD', PASSWORD), patch.object(app, 'initialize', fake_initialize):
            with TestClient(app.app) as client:
                self.assertEqual(client.get('/api/results').status_code, 401)
                board = client.post('/api/boards', headers=HEADERS, json={'legs': legs})
                self.assertEqual(board.status_code, 200, board.text)
                self.assertEqual(client.get('/api/results', headers=HEADERS).json()['entries'], 0)
                entry = client.post('/api/entries', headers=HEADERS,
                    json={'legs': legs, 'stake': 10, 'entry_type': 'power', 'payout_multiplier': 3}).json()
                master = app.app.state.master
                master.upsert_game({'id': 'test-game', 'date': '2026-10-27', 'status': 'final',
                                    'home_team': 'NYK', 'away_team': 'BOS', 'source': 'test'})
                for player, team, stat, value in [('Jalen Brunson', 'NYK', 'pts', 24), ('Derrick White', 'BOS', 'ast', 5)]:
                    master.upsert_player_game({'game_id': 'test-game', 'player': player, 'team': team,
                        'date': '2026-10-27', 'appearance_state': 'Played', stat: value, 'source': 'test'})
                master.grade_entries()
                self.assertEqual(master.get_entry(entry['id'])['provisional_status'], 'won')
                self.assertIsNone(master.get_entry(entry['id'])['actual_payout'])
                response = client.post('/api/entries/'+entry['id']+'/settle', headers=HEADERS,
                                       json={'payout':30,'note':'Actual PrizePicks slip checked'})
                self.assertEqual(response.status_code, 200, response.text)
                self.assertEqual(client.get('/api/results', headers=HEADERS).json()['net_profit'], 20)
                export = client.get('/api/export/master.xlsx', headers=HEADERS)
                workbook = openpyxl.load_workbook(BytesIO(export.content), read_only=True)
                self.assertIn('entries', workbook.sheetnames)
                workbook.close()
                backup = client.get('/api/export/backup.sqlite', headers=HEADERS)
                backup_path = self.root / 'restored.sqlite'
                backup_path.write_bytes(backup.content)
                with sqlite3.connect(backup_path) as db:
                    self.assertEqual(db.execute('SELECT actual_payout FROM entries').fetchone()[0], 30)
            with TestClient(app.app) as client:
                summary = client.get('/api/results', headers=HEADERS).json()
                self.assertEqual(summary['settled_entries'], 1)
                self.assertEqual(summary['net_profit'], 20)

    def test_export_isolation_and_formula_escape(self):
        master = MasterStore(self.root/'master.sqlite')
        master.record_news({'id':'n1','title':'=HYPERLINK("https://untrusted.example")','source':'test'})
        with ThreadPoolExecutor(max_workers=3) as pool:
            paths = list(pool.map(lambda n: app.export_master(master, self.root/'exports'), range(3)))
        for path in paths:
            workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
            row = list(workbook['news'].values)[1]
            self.assertTrue(str(row[1]).startswith("'="))
            workbook.close()
        self.assertFalse(list((self.root/'exports').glob('*.partial.xlsx')))

    def test_feed_failure_alerts_are_transition_based(self):
        master = MasterStore(self.root/'master.sqlite')
        messages=[]
        async def notify(text,event_key=None): messages.append((text,event_key))
        async def run():
            master.feed_failure('espn_news','HTTPStatusError')
            await app.notify_feed_health(master,notify)
            await app.notify_feed_health(master,notify)
            self.assertEqual(len(messages),1)
            master.feed_success('espn_news','Recovered')
            await app.notify_feed_health(master,notify)
            self.assertEqual(len(messages),2)
        asyncio.run(run())
        self.assertIn('healthy',messages[-1][0])

    def test_model_initialization_retries_after_failure(self):
        calls=[]
        class State: phase='starting'
        class Application: state=State()
        async def initialize(application):
            calls.append(1);application.state.phase='error' if len(calls)==1 else 'ready'
        async def pause(seconds): pass
        with patch.object(app,'initialize',initialize),patch.object(app.asyncio,'sleep',pause):
            asyncio.run(app.initialize_with_retry(Application()))
        self.assertEqual(len(calls),2)


    def test_manual_csv_import_updates_master_and_grades_without_claiming_payout(self):
        with patch.object(app,'DATA_DIR',self.root),patch.object(app,'CHAT_PASSWORD',PASSWORD),patch.object(app,'initialize',fake_initialize):
            with TestClient(app.app) as client:
                legs=[{'player':'Jalen Brunson','team':'NYK','stat':'pts','side':'more','line':20.5,'date':'2026-10-27'},
                      {'player':'Cade Cunningham','team':'DET','stat':'pts','side':'less','line':30.5,'date':'2026-10-27'}]
                entry=client.post('/api/entries',headers=HEADERS,json={'legs':legs,'stake':10,'entry_type':'power','payout_multiplier':3}).json()
                content='date,player,team,opponent,game_status,appearance_state,minutes,pts\n2026-10-27,Jalen Brunson,NYK,DET,final,Played,35:30,24\n2026-10-27,Cade Cunningham,DET,NYK,final,Played,36,25\n'
                response=client.post('/api/import/stats',headers=HEADERS,json={'filename':'finals.csv','content':content})
                self.assertEqual(response.status_code,200,response.text)
                self.assertEqual(response.json()['status'],'complete')
                saved=app.app.state.master.get_entry(entry['id'])
                self.assertEqual(saved['provisional_status'],'won')
                self.assertIsNone(saved['actual_payout'])
                self.assertIn('Manual',saved['grades'][0]['source'])
                wrong=client.post('/api/import/stats',headers=HEADERS,json={'filename':'board.csv','content':'player,stat,line,side\nJalen Brunson,pts,22.5,more\n'})
                self.assertEqual(wrong.status_code,422)
                self.assertEqual(app.app.state.master.status()['counts']['player_games'],2)


if __name__ == '__main__': unittest.main()

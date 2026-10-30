"""One-time import of the embedded workbook into the durable NBA master.

Game and player keys match the live ESPN adapter. The import supplements missing
history; it never overwrites an existing live observation or reimports on restart.
The workbook is an embedded subset, not the unavailable full companion archive.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import math
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

from feeds import canonical_game, canonical_player, canonical_team
from master import GAME_FIELDS, PLAYER_FIELDS

STAT_FIELDS = ('pts', 'reb', 'ast', 'tpm', 'fga', 'tpa', 'minutes', 'stl', 'blk', 'tov', 'fgm','ftm','fta','oreb','dreb','fouls')


def _date(value):
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    try:
        return dt.date.fromisoformat(str(value)[:10]).isoformat()
    except (TypeError, ValueError):
        return None


def _number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (ValueError, TypeError):
        return None
    return number if math.isfinite(number) and number >= 0 else None


def seed_workbook(store, path):
    """Import each SHA once and never replace already-collected live statistics."""
    workbook_path = Path(path)
    sha = hashlib.sha256(workbook_path.read_bytes()).hexdigest()
    with store.connect() as db:
        db.execute('''CREATE TABLE IF NOT EXISTS seed_imports (
            sha256 TEXT PRIMARY KEY, filename TEXT NOT NULL, completed_at TEXT NOT NULL,
            games_added INTEGER NOT NULL, player_games_added INTEGER NOT NULL,
            observed_rows INTEGER NOT NULL, skipped_rows INTEGER NOT NULL)''')
        old = db.execute('SELECT * FROM seed_imports WHERE sha256=?', (sha,)).fetchone()
        if old:
            return {'imported': False, 'skipped': True, 'sha256': sha, 'games': 0,
                    'player_games': 0, 'original_import': dict(old),
                    'coverage': 'Embedded workbook subset only; full companion ZIP not imported.'}
        existing_games = {row[0] for row in db.execute('SELECT id FROM games')}
        existing_players = {(row[0], row[1]) for row in
                            db.execute('SELECT game_id,player_id FROM player_games')}
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    games, observations = {}, []
    observed, skipped = 0, 0
    today = dt.datetime.now(ZoneInfo('America/New_York')).date().isoformat()
    source = f'workbook:{workbook_path.name}#sha256={sha};sheet=Player_Games'
    try:
        if 'Player_Games' not in workbook.sheetnames:
            raise ValueError('Workbook has no Player_Games sheet.')
        sheet = workbook['Player_Games']
        headers = list(next(sheet.iter_rows(min_row=7, max_row=7, values_only=True)))
        required = {'player', 'date', 'team', 'opponent'}
        if not required.issubset(headers):
            raise ValueError('Player_Games row 7 does not contain required headers.')
        for excel_row, values in enumerate(sheet.iter_rows(min_row=8, values_only=True), 8):
            row = dict(zip(headers, values))
            if not any(v is not None for v in values):
                # Other worksheet sections can have different grains/columns.
                # Only the contiguous canonical table under row 7 is seed data.
                if observed:
                    break
                continue
            observed += 1
            date = _date(row.get('date'))
            player = str(row.get('player') or '').strip()
            team, opponent = canonical_team(row.get('team')), canonical_team(row.get('opponent'))
            if not date or not player or not team or not opponent or team == opponent:
                skipped += 1
                continue
            # A completed snapshot cannot establish participation for future records.
            if date > today:
                skipped += 1
                continue
            game_id = canonical_game(date, team, opponent)
            player_id = canonical_player(player)
            state = str(row.get('appearance_state') or 'Unknown')
            venue = str(row.get('venue') or '').casefold()
            home = team if venue == 'home' else opponent if venue == 'away' else None
            away = opponent if venue == 'home' else team if venue == 'away' else None
            if game_id not in games:
                games[game_id] = {'id': game_id, 'date': date,
                                  'status': 'final' if state.casefold() == 'played' else 'unknown',
                                  'source': source, 'provider_game_id': row.get('game_id')}
                if home and away:
                    games[game_id].update(home_team=home, away_team=away)
            elif state.casefold() == 'played':
                games[game_id]['status'] = 'final'
            if home and away:
                games[game_id].update(home_team=home, away_team=away)
            if (game_id, player_id) in existing_players:
                continue
            stat = {key: _number(row.get(key)) for key in STAT_FIELDS}
            stat.update(game_id=game_id, player_id=player_id, player=player, team=team,
                        opponent=opponent, date=date, season=row.get('season'),
                        season_type=row.get('season_type'), appearance_state=state,
                        source=source + f';row={excel_row}',
                        original_game_id=row.get('game_id'),
                        original_player_id=row.get('player_id'),
                        source_ids=row.get('source_ids'))
            observations.append(stat)
            existing_players.add((game_id, player_id))
    finally:
        workbook.close()
    completed_at = dt.datetime.now(dt.timezone.utc).isoformat()
    games_added = player_games_added = 0
    with store.connect() as db:
        # INSERT OR IGNORE is essential: live collectors can insert/correct the same
        # event while the workbook is being parsed. A prior existence check alone
        # would race and a normal upsert could overwrite those newer observations.
        for game_id, game in games.items():
            if game_id in existing_games:
                continue
            fields = GAME_FIELDS + ['updated_at']
            values = [game.get(key) for key in GAME_FIELDS] + [completed_at]
            cursor = db.execute('INSERT OR IGNORE INTO games (' + ','.join(fields) + ') VALUES ('
                                + ','.join('?' for _ in fields) + ')', values)
            if cursor.rowcount:
                games_added += 1
                store._snapshot(db, 'games', game_id, game)
        for stat in observations:
            fields = PLAYER_FIELDS + ['updated_at']
            values = [stat.get(key) for key in PLAYER_FIELDS] + [completed_at]
            cursor = db.execute('INSERT OR IGNORE INTO player_games (' + ','.join(fields) + ') VALUES ('
                                + ','.join('?' for _ in fields) + ')', values)
            if cursor.rowcount:
                player_games_added += 1
                store._snapshot(db, 'player_games', stat['game_id'] + ':' + stat['player_id'], stat)
        db.execute('INSERT OR IGNORE INTO seed_imports VALUES (?,?,?,?,?,?,?)',
                   (sha, workbook_path.name, completed_at, games_added,
                    player_games_added, observed, skipped))
    store.feed_success('workbook_seed', f'Imported embedded snapshot: {player_games_added} observations; '
                       'not a live feed and not the full companion archive.')
    return {'imported': True, 'skipped': False, 'sha256': sha, 'games': games_added,
            'player_games': player_games_added, 'observed_rows': observed, 'skipped_rows': skipped,
            'coverage': 'Embedded workbook subset only; full companion ZIP not imported.'}

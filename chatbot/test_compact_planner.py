"""Bounded model retrieval still reaches exact workbook rows and scoped tracking."""
import asyncio
from contextlib import ExitStack
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

import httpx
import openpyxl

import app
from master import MasterStore
from workbook import WorkbookStore


class CompactPlannerTests(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        root = Path(self.directory.name)
        book = openpyxl.Workbook()
        book.active.title = 'Jalen Brunson'
        book.active.append(['Player summary'])
        book.active.append(['Potential assists', 12.5])
        games = book.create_sheet('Player_Games')
        for _ in range(6):
            games.append([None])
        games.append(['player', 'pts'])
        games.append(['Jalen Brunson', 20])
        for index in range(50):
            book.create_sheet('Unrelated ' + str(index)).append(['private irrelevant snapshot'])
        book.save(root / 'source.xlsx')
        book.close()
        self.store = WorkbookStore(str(root / 'source.xlsx'), str(root / 'workbook.sqlite'))
        self.master = MasterStore(str(root / 'master.sqlite'))
        with self.master.connect() as db:
            for season, entity, value in [('2026-27', 'Jalen Brunson', 12.5),
                                           ('2025-26', 'Jalen Brunson', 99),
                                           ('2026-27', 'Other Player', 88)]:
                db.execute('INSERT INTO refreshed_tracking VALUES (?,?,?,?)',
                    (season, entity, 'Passing', json.dumps({'gp': 5,
                     'fields': {'potential_assists': value}, 'data_through_date': '2026-10-28'})))

    def tearDown(self):
        self.directory.cleanup()

    def routes(self, question='Read the stored passing record for Jalen Brunson'):
        stack = ExitStack()
        resolved = {'question': question,
                    'players': ['Jalen Brunson'], 'teams': []}
        stack.enter_context(patch('basketball_research.resolve_question', return_value=resolved))
        stack.enter_context(patch('basketball_research.answer_assignment_question', return_value=None))
        stack.enter_context(patch('basketball_research.canonical_history', return_value=[]))
        stack.enter_context(patch('basketball_research.research_context', return_value={}))
        stack.enter_context(patch('matchup_queries.answer_matchup_question', return_value=None))
        stack.enter_context(patch('player_stats.answer_player_question', return_value=None))
        stack.enter_context(patch('grounded_reply.evidence_cards', return_value=[]))
        return stack

    async def run_tool(self, name, arguments):
        seen = []

        def response(request):
            payload = json.loads(request.content)
            seen.append(payload)
            if len(seen) == 1:
                message = {'content': '', 'tool_calls': [{'function': {
                    'name': name, 'arguments': arguments}}]}
            else:
                message = {'content': 'The requested records have been retrieved.'}
            return httpx.Response(200, json={'message': message})

        real_client = httpx.AsyncClient
        final = AsyncMock(return_value={'answer': 'Evidence reviewed', 'sources': []})
        with self.routes(), patch.object(app, 'finish_retrieved_reply', final), patch.object(
                app.httpx, 'AsyncClient', side_effect=lambda **kwargs: real_client(
                    transport=httpx.MockTransport(response), **kwargs)):
            await app.answer_question(self.store, app.ChatRequest(message='Read the record'), {}, self.master)
        return seen, final.await_args.args[-1]

    async def test_original_workbook_row_remains_accessible_with_compact_catalog(self):
        seen, cards = await self.run_tool('read_sheet',
            {'sheet': 'Jalen Brunson', 'start_row': 2, 'limit': 1})
        tool_result = json.loads(seen[1]['messages'][-1]['content'])
        self.assertEqual(tool_result['rows'], [{'_row': 2, 'A': 'Potential assists', 'B': 12.5}])
        self.assertEqual(cards[0]['kind'], 'workbook_excerpt')
        self.assertIn('Potential assists', cards[0]['text'])
        self.assertNotIn('Unrelated 49', seen[0]['messages'][0]['content'])
        self.assertEqual(seen[0]['options']['num_ctx'], 4096)
        self.assertEqual(len(seen[0]['tools']), 8)
        self.assertEqual(len(seen), 2)

    async def test_tracking_query_keeps_requested_player_season_and_metric(self):
        sql = """SELECT entity, season,
            json_extract(payload_json,'$.fields.potential_assists') AS potential_assists,
            json_extract(payload_json,'$.gp') AS games,
            json_extract(payload_json,'$.data_through_date') AS through_date
            FROM refreshed_tracking WHERE entity='Jalen Brunson'
            AND season='2026-27' AND category='Passing'"""
        seen, cards = await self.run_tool('query_master', {'sql': sql})
        result = json.loads(seen[1]['messages'][-1]['content'])
        self.assertEqual(result['rows'], [['Jalen Brunson', '2026-27', 12.5, 5, '2026-10-28']])
        self.assertEqual(cards[0]['query'], sql)
        self.assertIn('potential_assists', cards[0]['text'])
        self.assertIn('filters and calculation need review', cards[0]['text'])

    async def test_planner_timeout_stops_without_retry_and_keeps_truthful_failure(self):
        calls = []

        async def response(request):
            calls.append(request)
            await asyncio.sleep(10)

        real_client = httpx.AsyncClient
        final = AsyncMock()
        with self.routes(), patch.object(app, 'PLANNER_TIMEOUT_SECONDS', .03), patch.object(
                app, 'finish_retrieved_reply', final), patch.object(app.httpx, 'AsyncClient',
                side_effect=lambda **kwargs: real_client(transport=httpx.MockTransport(response), **kwargs)):
            result = await app.answer_question(self.store, app.ChatRequest(message='Read the record'), {}, self.master)
        self.assertEqual(len(calls), 1)
        final.assert_not_awaited()
        self.assertEqual(result['ai_failure_category'], 'timeout')
        self.assertFalse(result['ai_used'])

    async def test_repeated_tool_requests_stop_after_two_rounds(self):
        calls = []

        def response(request):
            calls.append(request)
            return httpx.Response(200, json={'message': {'content': '', 'tool_calls': [
                {'function': {'name': 'read_sheet', 'arguments':
                    {'sheet': 'Jalen Brunson', 'start_row': 2, 'limit': 1}}}]}})

        real_client = httpx.AsyncClient
        final = AsyncMock(return_value={'answer': 'Reviewed', 'sources': []})
        with self.routes(), patch.object(app, 'finish_retrieved_reply', final), patch.object(
                app.httpx, 'AsyncClient', side_effect=lambda **kwargs: real_client(
                    transport=httpx.MockTransport(response), **kwargs)):
            await app.answer_question(self.store, app.ChatRequest(message='Read the record'), {}, self.master)
        self.assertEqual(len(calls), 2)
        final.assert_awaited_once()
        self.assertTrue(all(card['kind'] == 'workbook_excerpt' for card in final.await_args.args[-1]))

    async def test_final_reply_obeys_remaining_retrieval_deadline(self):
        async def slow_final(*args):
            await asyncio.sleep(10)

        real_client = httpx.AsyncClient
        response = httpx.MockTransport(lambda request: httpx.Response(200,
            json={'message': {'content': 'No further retrieval needed.'}}))
        with self.routes(), patch.object(app, 'RETRIEVAL_REPLY_TIMEOUT_SECONDS', .03), patch.object(
                app, 'finish_retrieved_reply', slow_final), patch.object(app.httpx, 'AsyncClient',
                side_effect=lambda **kwargs: real_client(transport=response, **kwargs)):
            result = await app.answer_question(self.store, app.ChatRequest(message='Read the record'), {}, self.master)
        self.assertEqual(result['ai_failure_category'], 'timeout')
        self.assertFalse(result['ai_used'])

    async def test_parallel_tool_previews_share_one_character_budget(self):
        seen = []

        def response(request):
            payload = json.loads(request.content)
            seen.append(payload)
            message = {'content': 'Retrieved the requested excerpts.'}
            if len(seen) == 1:
                message = {'content': '', 'tool_calls': [{'function': {'name': 'read_sheet',
                    'arguments': {'sheet': 'Jalen Brunson', 'start_row': index + 1}}}
                    for index in range(4)]}
            return httpx.Response(200, json={'message': message})

        real_client = httpx.AsyncClient
        final = AsyncMock(return_value={'answer': 'Reviewed', 'sources': []})
        large_row = {'rows': [{'_row': 1, 'A': ('"é\\' * 1500)}]}
        with self.routes(), patch.object(self.store, 'read', return_value=large_row), patch.object(
                app, 'finish_retrieved_reply', final), patch.object(app.httpx, 'AsyncClient',
                side_effect=lambda **kwargs: real_client(transport=httpx.MockTransport(response), **kwargs)):
            await app.answer_question(self.store, app.ChatRequest(message='Read the record'), {}, self.master)
        previews = [message['content'] for message in seen[1]['messages'] if message['role'] == 'tool']
        self.assertEqual(len(previews), 4)
        self.assertLessEqual(sum(map(len, previews)), 3600)
        self.assertTrue(all(len(text) <= 900 for text in previews))
        self.assertTrue(all(json.loads(text)['truncated'] for text in previews))

    async def test_named_player_row_goes_directly_to_original_workbook_tool(self):
        question = 'Read row 2 of the Jalen Brunson worksheet, rather than his overall averages.'
        seen = []

        def response(request):
            payload = json.loads(request.content)
            seen.append(payload)
            message = {'content': 'The original source row is now available.'}
            if len(seen) == 1:
                message = {'content': '', 'tool_calls': [{'function': {'name': 'read_sheet',
                    'arguments': {'sheet': 'Jalen Brunson', 'start_row': 2, 'limit': 1}}}]}
            return httpx.Response(200, json={'message': message})

        real_client = httpx.AsyncClient
        final = AsyncMock(return_value={'answer': 'Read the original row', 'sources': []})
        with self.routes(question) as stack:
            for target in ('basketball_research.answer_assignment_question',
                           'basketball_research.research_context',
                           'matchup_queries.answer_matchup_question', 'player_stats.answer_player_question'):
                stack.enter_context(patch(target, side_effect=AssertionError('Irrelevant baseline route')))
            stack.enter_context(patch.object(app, 'generated_evidence_reply',
                AsyncMock(side_effect=AssertionError('Must retrieve the source first'))))
            stack.enter_context(patch.object(app, 'finish_retrieved_reply', final))
            stack.enter_context(patch.object(app.httpx, 'AsyncClient',
                side_effect=lambda **kwargs: real_client(transport=httpx.MockTransport(response), **kwargs)))
            result = await app.answer_question(self.store, app.ChatRequest(message=question), {}, self.master)
        self.assertIn('Potential assists', result['answer'])
        self.assertIn('12.5', result['answer'])
        self.assertIn('"_row": 2', result['answer'])
        final.assert_not_awaited()
        self.assertEqual(len(seen), 1)
        self.assertEqual(seen[0]['messages'][-1]['content'], app.final_answer_prompt(question))

    async def test_telegram_source_request_skips_early_statistics_and_assignment_routes(self):
        question = 'Show cell B2 in the Jalen Brunson worksheet.'
        state = SimpleNamespace(store=self.store, master=self.master, phase='ready',
                                detail='Ready', chat_lock=asyncio.Lock())
        answer = AsyncMock(return_value={'answer': 'Original source row', 'ai_used': False})
        with self.routes(question) as stack:
            for target in ('basketball_research.answer_assignment_question',
                           'matchup_queries.answer_matchup_question', 'player_stats.answer_player_question'):
                stack.enter_context(patch(target, side_effect=AssertionError('Irrelevant baseline route')))
            stack.enter_context(patch.object(app.app, 'state', state))
            stack.enter_context(patch.object(app, 'answer_question', answer))
            result = await app.telegram_answer(question)
        self.assertEqual(result, 'Original source row')
        answer.assert_awaited_once()
        self.assertFalse(state.chat_lock.locked())

    def test_source_location_routing_preserves_ordinary_metric_questions(self):
        for question in ('Read row 12 of Jalen Brunson', 'Inspect the Matchup_Assignments sheet',
                         'Show cell B12 of that worksheet', 'Read range A12:C14', 'Open the Overview tab'):
            self.assertTrue(app.explicit_source_request(question), question)
        for question in ('What is Jalen Brunson average potential assists in the stored passing sheet?',
                         'Who is his matchup?', 'Show his average points', 'What happened in his last two games?'):
            self.assertFalse(app.explicit_source_request(question), question)

    def test_only_original_positive_integer_row_numbers_are_reference_metadata(self):
        for tool, key in [('read_sheet', '_row'), ('search_workbook', 'row')]:
            result_key = 'rows' if tool == 'read_sheet' else 'matches'
            valid = app.query_evidence_card(tool, {}, {result_key: [{key: 12, 'A': 'Example'}]}, 1)
            self.assertEqual(valid['reference_rows'], [12])
            for invalid in (True, '12', 0, -1, 1.5):
                card = app.query_evidence_card(tool, {}, {result_key: [{key: invalid}]}, 1)
                self.assertEqual(card['reference_rows'], [])
        sql = app.query_evidence_card('query_workbook', {'sql': 'SELECT 12 AS row'},
                                     {'columns': ['row'], 'rows': [[12]]}, 1)
        self.assertNotIn('reference_rows', sql)

    def test_context_bounds_catalog_but_preserves_master_query_fields(self):
        context = app.compact_chat_context(self.store, self.master,
                    {'players': ['Jalen Brunson'], 'teams': []})
        value = json.loads(context)
        self.assertLessEqual(len(context), 2300)
        self.assertEqual([sheet['sheet'] for sheet in value['workbook_catalog']],
                         ['Jalen Brunson', 'Player_Games'])
        self.assertIn('appearance_state', value['master_tables']['player_games'])
        self.assertIn('payload_json', value['master_tables']['refreshed_tracking'])
        self.assertIn('fields', value['tracking_json_keys'])
        self.assertIn('search_workbook', value['other_sheets'])


if __name__ == '__main__':
    unittest.main()

"""Read cached Excel values into a read-only, queryable workbook snapshot."""
from __future__ import annotations

import datetime as dt
import json
import re
import sqlite3
import time
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter


def quote(identifier: str) -> str:
    return '"' + identifier.replace('"', '""') + '"'


def cell_value(value):
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float)):
        return value
    return str(value)


class WorkbookStore:
    def __init__(self, workbook_path: str, database_path: str):
        self.database_path = str(Path(database_path).resolve())
        Path(self.database_path).parent.mkdir(parents=True, exist_ok=True)
        self.sheets = {}
        workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
        db = sqlite3.connect(self.database_path)
        try:
            db.execute('CREATE TABLE search_rows (sheet TEXT, excel_row INTEGER, text TEXT)')
            for number, sheet in enumerate(workbook):
                raw_table = f"raw_{number:02d}"
                rows = []
                width = 1
                for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    cells = [cell_value(v) for v in row]
                    while cells and cells[-1] is None:
                        cells.pop()
                    if not cells:
                        continue
                    width = max(width, len(cells))
                    rows.append((row_number, cells))
                letters = [get_column_letter(i + 1) for i in range(width)]
                # No affinity: keep Excel numbers numeric, missing cells NULL, IDs text.
                db.execute(f'CREATE TABLE {quote(raw_table)} (_row INTEGER PRIMARY KEY, '
                           + ', '.join(quote(c) for c in letters) + ')')
                db.executemany(f'INSERT INTO {quote(raw_table)} VALUES ('
                               + ','.join('?' for _ in range(width + 1)) + ')',
                               [(n, *cells, *([None] * (width - len(cells)))) for n, cells in rows])
                db.executemany('INSERT INTO search_rows VALUES (?,?,?)',
                               [(sheet.title, n, ' | '.join(str(v) for v in cells if v is not None))
                                for n, cells in rows])
                header = next((cells for n, cells in rows if n == 7), [])
                structured = (len(header) >= 3 and all(isinstance(v, str) and v for v in header)
                              and len(set(header)) == len(header))
                fields = []
                if structured:
                    fields = [re.sub(r'[^a-z0-9]+', '_', v.lower()).strip('_') for v in header]
                    structured = len(set(fields)) == len(fields) and all(fields)
                if structured:
                    # Named views cover the first contiguous table only. Later sections
                    # can have different columns; they stay accessible via raw Excel rows.
                    table_end = 7
                    for n, cells in rows:
                        if n <= 7:
                            continue
                        if n != table_end + 1 or cells == header:
                            break
                        table_end = n
                    aliases = ','.join(f'{quote(letters[i])} AS {quote(name)}'
                                       for i, name in enumerate(fields))
                    db.execute(f'CREATE VIEW {quote(sheet.title)} AS SELECT _row,{aliases} '
                               f'FROM {quote(raw_table)} WHERE _row BETWEEN 8 AND {table_end}')
                self.sheets[sheet.title] = {
                    'sheet': sheet.title, 'last_row': rows[-1][0] if rows else 0,
                    'nonempty_rows': len(rows), 'raw_table': raw_table,
                    'sql_table': sheet.title if structured else raw_table,
                    'columns': ['_row'] + (fields if structured else letters),
                    'named_table_rows': [8, table_end] if structured else None,
                    'notes': [{'row': n, 'text': str(cells[0])[:1400]}
                              for n, cells in rows if n <= 3 and cells],
                }
            db.commit()
        finally:
            workbook.close()
            db.close()

    def connect(self):
        return sqlite3.connect(Path(self.database_path).as_uri() + '?mode=ro', uri=True)

    def describe(self, sheet: str):
        if sheet not in self.sheets:
            raise ValueError('Unknown sheet. Available sheets: ' + ', '.join(self.sheets))
        return self.sheets[sheet]

    def read(self, sheet: str, start_row: int = 1, limit: int = 12):
        info = self.describe(sheet)
        start_row = max(1, int(start_row))
        limit = min(40, max(1, int(limit)))
        db = self.connect()
        try:
            cursor = db.execute(f'SELECT * FROM {quote(info["raw_table"])} '
                                'WHERE _row >= ? AND _row < ? ORDER BY _row',
                                (start_row, start_row + limit))
            names = [c[0] for c in cursor.description]
            return {'sheet': sheet, 'rows': [dict((k, v) for k, v in zip(names, row)
                                                  if v is not None) for row in cursor],
                    'last_row': info['last_row']}
        finally:
            db.close()

    def search(self, text: str, sheet: str = ''):
        if not text or len(text) > 160:
            raise ValueError('Search text must contain 1–160 characters.')
        if sheet:
            self.describe(sheet)
        db = self.connect()
        try:
            escaped = text.replace('\\', '\\\\').replace('%', '\\%').replace('_', '\\_')
            sql = "SELECT sheet,excel_row,text FROM search_rows WHERE text LIKE ? ESCAPE '\\'"
            params = ['%' + escaped + '%']
            if sheet:
                sql += ' AND sheet = ?'
                params.append(sheet)
            return {'matches': [{'sheet': s, 'row': n, 'text': t[:1600]}
                                for s, n, t in db.execute(sql + ' LIMIT 8', params)],
                    'note': 'At most 8 matches; use a sheet filter to narrow the search.'}
        finally:
            db.close()

    def query(self, sql: str):
        if not isinstance(sql, str) or not 1 <= len(sql) <= 5000:
            raise ValueError('Use a SELECT query of at most 5000 characters.')
        db = self.connect()
        deadline = time.monotonic() + 2
        allowed = {sqlite3.SQLITE_SELECT, sqlite3.SQLITE_READ, sqlite3.SQLITE_FUNCTION,
                   sqlite3.SQLITE_RECURSIVE}
        blocked_functions = {'load_extension', 'readfile', 'writefile', 'randomblob',
                             'zeroblob', 'printf', 'format'}

        def authorize(action, arg1, arg2, database, trigger):
            if action not in allowed:
                return sqlite3.SQLITE_DENY
            if action == sqlite3.SQLITE_FUNCTION and str(arg2).lower() in blocked_functions:
                return sqlite3.SQLITE_DENY
            return sqlite3.SQLITE_OK

        db.set_authorizer(authorize)
        db.set_progress_handler(lambda: int(time.monotonic() > deadline), 1000)
        db.setlimit(sqlite3.SQLITE_LIMIT_LENGTH, 100_000)
        db.setlimit(sqlite3.SQLITE_LIMIT_EXPR_DEPTH, 50)
        try:
            cursor = db.execute(sql)
            if cursor.description is None:
                raise ValueError('Only read-only queries are supported.')
            rows = cursor.fetchmany(61)
            return {'columns': [c[0] for c in cursor.description], 'rows': rows[:60],
                    'truncated': len(rows) > 60,
                    'note': 'Counts and averages cover the SQL filter, not just the returned rows. '
                            'Cite the sheet and filter; _row contains the original Excel row number.'}
        finally:
            db.close()

    def context(self):
        important = ['Player_Games', 'Calculated_Splits', 'Pair_Tracking', 'Pair_Summaries',
                     'Team_Schedule', 'Season_Totals']
        catalog = [{'sheet': name, 'sql_table': info['sql_table'], 'last_row': info['last_row']}
                   for name, info in self.sheets.items()]
        schema = {name: self.sheets[name]['columns'] for name in important if name in self.sheets}
        return json.dumps({'sheets': catalog, 'schema': schema,
                           'overview': self.read('Overview', 1, 20)}, default=str)

import datetime as dt
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

import openpyxl

from analysis_engine import analyze_board, analysis_history
from feeds import canonical_game, canonical_player
from master import MasterStore
from seed_master import seed_workbook


class SeedAndAnalysisTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.root = Path(self.directory.name)
        self.store = MasterStore(self.root / 'master.sqlite')

    def tearDown(self):
        self.directory.cleanup()

    def workbook(self, rows, name='seed.xlsx'):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = 'Player_Games'
        for _ in range(6):
            sheet.append([None])
        sheet.append(['player', 'date', 'team', 'opponent', 'venue', 'appearance_state',
                      'pts', 'reb', 'ast', 'season', 'season_type', 'game_id', 'player_id'])
        for row in rows:
            sheet.append(row)
        path = self.root / name
        workbook.save(path)
        workbook.close()
        return path

    def observation(self, day, player='Jalen Brunson', **changes):
        date = f'2025-01-{day:02d}'
        game_id = canonical_game(date, 'NYK', 'BOS')
        self.store.upsert_game({'id': game_id, 'date': date, 'home_team': 'NYK',
                               'away_team': 'BOS', 'status': 'final', 'source': 'fixture'})
        record = {'game_id': game_id, 'player_id': canonical_player(player), 'player': player,
                  'team': 'NYK', 'opponent': 'BOS', 'date': date, 'season': '2024-25',
                  'season_type': 'Regular Season', 'appearance_state': 'Played',
                  'pts': day, 'reb': 3, 'ast': 2, 'source': 'fixture'}
        record.update(changes)
        self.store.upsert_player_game(record)
        return record

    def test_seed_is_once_only_and_never_overwrites_collected_stats(self):
        game = canonical_game('2025-01-01', 'NYK', 'BOS')
        self.observation(1, pts=99, source='live-provider')
        rows = [
            ['Jalen Brunson', dt.datetime(2025, 1, 1, 12), 'NY', 'BOS', 'Home', 'Played',
             20, 3, 4, '2024-25', 'Regular Season', 'official-id-1', 'official-player-id'],
            ['Jalen Brunson', dt.datetime(2025, 1, 2, 12), 'NYK', 'BOS', 'Away', 'Played',
             24, None, 5, '2024-25', 'Regular Season', 'official-id-2', 'official-player-id'],
        ]
        path = self.workbook(rows)
        first = seed_workbook(self.store, path)
        self.assertTrue(first['imported'])
        self.assertEqual(first['player_games'], 1)
        self.assertEqual(first['games'], 1)
        second = seed_workbook(MasterStore(self.store.path), path)
        self.assertTrue(second['skipped'])
        with self.store.connect() as db:
            old = db.execute('SELECT pts,source FROM player_games WHERE game_id=?', (game,)).fetchone()
            self.assertEqual(tuple(old), (99.0, 'live-provider'))
            row = db.execute('SELECT reb,source FROM player_games WHERE date=?', ('2025-01-02',)).fetchone()
            self.assertIsNone(row[0])
            self.assertIn('sheet=Player_Games;row=9', row[1])
            self.assertEqual(db.execute('SELECT COUNT(*) FROM seed_imports').fetchone()[0], 1)

    def test_new_workbook_hash_only_supplements_new_rows(self):
        rows = [['Jalen Brunson', dt.datetime(2025, 1, 1), 'NYK', 'BOS', 'Home', 'Played',
                 20, 3, 4, '2024-25', 'Regular Season']]
        seed_workbook(self.store, self.workbook(rows))
        self.observation(1, pts=42, source='corrected-live')
        rows.append(['Josh Hart', dt.datetime(2025, 1, 1), 'NYK', 'BOS', 'Home', 'Played',
                     10, 8, 3, '2024-25', 'Regular Season'])
        revised = seed_workbook(self.store, self.workbook(rows, 'revised.xlsx'))
        self.assertEqual(revised['player_games'], 1)
        with self.store.connect() as db:
            self.assertEqual(db.execute('SELECT pts FROM player_games WHERE player=?',
                                        ('Jalen Brunson',)).fetchone()[0], 42)

    def test_live_insert_during_workbook_parse_wins_the_seed_race(self):
        path = self.workbook([['Jalen Brunson', dt.datetime(2025, 1, 1), 'NYK', 'BOS',
                               'Home', 'Played', 20, 3, 4, '2024-25', 'Regular Season']])
        loader = openpyxl.load_workbook

        def insert_live_before_parse(*args, **kwargs):
            self.observation(1, pts=99, source='live-arrived-during-seed')
            return loader(*args, **kwargs)

        with patch('seed_master.openpyxl.load_workbook', side_effect=insert_live_before_parse):
            result = seed_workbook(self.store, path)
        self.assertEqual(result['games'], 0)
        self.assertEqual(result['player_games'], 0)
        with self.store.connect() as db:
            row = db.execute('SELECT pts,source FROM player_games').fetchone()
            self.assertEqual(tuple(row), (99.0, 'live-arrived-during-seed'))

    def test_seed_dnp_does_not_alone_establish_final_game(self):
        path = self.workbook([['Josh Hart', dt.datetime(2025, 1, 1), 'NYK', 'BOS', 'Home',
                               'DNP', None, None, None, '2024-25', 'Regular Season']])
        seed_workbook(self.store, path)
        self.assertEqual(self.store.games('2025-01-01')[0]['status'], 'unknown')

    def test_analysis_excludes_target_future_live_and_missing_components(self):
        for day in range(1, 16):
            self.observation(day)
        self.observation(1, pts=0, reb=0, ast=0)
        self.observation(10, reb=None)
        self.observation(14, pts=None)
        self.store.upsert_game({'id': canonical_game('2025-01-15', 'NYK', 'BOS'), 'status': 'live'})
        self.observation(16, pts=999)
        self.observation(17, pts=999)
        self.store.upsert_game({'id': canonical_game('2025-01-16', 'NYK', 'BOS'), 'status': 'scheduled'})
        board = self.store.create_board([
            {'player': 'Jalen Brunson', 'stat': 'pts', 'line': 9.5, 'side': 'more', 'date': '2025-01-16'},
            {'player': 'Jalen Brunson', 'stat': 'pra', 'line': 20, 'side': 'less', 'date': '2025-01-16'},
        ])
        report = analyze_board(self.store, board['id'], as_of='2025-01-16T17:00:00Z')
        points, combo = report['legs']
        self.assertEqual(points['n'], 13)
        self.assertEqual(points['missing_games'], 1)
        self.assertEqual(points['more_count'], 4)
        self.assertEqual(points['sample_last_date'], '2025-01-13')
        self.assertEqual(combo['n'], 12)
        self.assertEqual(combo['missing_games'], 2)
        self.assertIn(0, [r['value'] for r in combo['sample_games']])
        self.assertNotIn(999, [r['value'] for r in points['sample_games']])
        self.assertTrue(points['eligible_for_ranking'])
        self.assertTrue(report['research_only'])
        self.assertEqual(self.store.get_board(board['id']), board)
        self.assertEqual(self.store.list_entries(), [])

    def test_analysis_small_sample_abstains_and_ties_remain_separate(self):
        for day in range(1, 6):
            self.observation(day, player='Josh Hart', pts=10)
        board = self.store.create_board([{'player': 'Josh Hart', 'stat': 'pts', 'line': 10,
                                         'side': 'more', 'date': '2025-01-10'}])
        report = analyze_board(self.store, board['id'])
        leg = report['legs'][0]
        self.assertEqual(leg['tie_count'], 5)
        self.assertEqual(leg['historical_selected_side_fraction'], 0)
        self.assertFalse(leg['eligible_for_ranking'])
        self.assertEqual(report['candidates'], [])

    def test_versions_persist_after_corrections_and_restart(self):
        for day in range(1, 13):
            self.observation(day)
        board = self.store.create_board([{'player': 'Jalen Brunson', 'stat': 'pts',
                                         'line': 10, 'side': 'more', 'date': '2025-01-20'}])
        first = analyze_board(self.store, board['id'])
        self.observation(12, pts=50)
        second = analyze_board(self.store, board['id'])
        self.assertNotEqual(first['id'], second['id'])
        history = analysis_history(MasterStore(self.store.path), board['id'])
        self.assertEqual(len(history), 2)
        by_id = {r['id']: r for r in history}
        self.assertEqual(by_id[first['id']]['legs'][0]['mean'], first['legs'][0]['mean'])
        self.assertNotEqual(first['legs'][0]['mean'], second['legs'][0]['mean'])

    def test_canonical_eight_knicks_names_and_accents_match(self):
        players = ['Jalen Brunson', 'Karl-Anthony Towns', 'Mikal Bridges', 'Josh Hart',
                   'Jose Alvarado', 'Andre Drummond', 'OG Anunoby', 'Miles McBride']
        self.assertEqual(len({canonical_player(p) for p in players}), 8)
        self.assertEqual(canonical_player('José Alvarado'), canonical_player('Jose Alvarado'))
        self.assertEqual(canonical_player('Karl Anthony Towns'), canonical_player('Karl-Anthony Towns'))
        self.assertNotEqual(canonical_player('Nic Claxton'), canonical_player('Nicolas Claxton'))

    def test_freshness_does_not_present_seed_as_live_or_official_coverage(self):
        self.store.feed_success('workbook_seed', 'static')
        board = self.store.create_board([{'player': 'Jalen Brunson', 'stat': 'pts',
                                         'line': 20, 'side': 'more', 'date': '2025-01-20'}])
        report = analyze_board(self.store, board['id'])
        self.assertEqual(report['freshness']['fresh_feeds'], [])
        self.assertTrue(any('No fresh availability' in w for w in report['freshness']['warnings']))
        self.store.feed_success('espn_availability', 'snapshots')
        report = analyze_board(self.store, board['id'])
        self.assertIn('espn_availability', report['freshness']['fresh_availability_feeds'])
        self.assertTrue(any('not complete official' in w for w in report['freshness']['warnings']))

    def test_expired_started_and_partial_period_boards_do_not_rank(self):
        for day in range(1, 13):
            self.observation(day)
        for extra in [
            {'date': '2025-01-13', 'period': '1H'},
            {'date': '2025-01-12'},
            {'date': '2025-01-16', 'team': 'NYK'},
        ]:
            self.store.upsert_game({'id': canonical_game('2025-01-16', 'NYK', 'BOS'),
                                   'date': '2025-01-16', 'home_team': 'NYK',
                                   'away_team': 'BOS', 'status': 'live'})
            leg = {'player': 'Jalen Brunson', 'stat': 'pts', 'line': 10, 'side': 'more', **extra}
            board = self.store.create_board([leg])
            result = analyze_board(self.store, board['id'], as_of='2025-01-16T17:00:00Z')
            self.assertEqual(result['candidates'], [])
            self.assertTrue(result['legs'][0]['eligibility_blocks'])

    def test_as_of_bounds_samples_and_ignores_maintenance_freshness(self):
        for day in range(1, 16):
            self.observation(day, pts=day * 10)
        self.store.feed_success('local_backup')
        board = self.store.create_board([{'player': 'Jalen Brunson', 'stat': 'pts',
                                         'line': 20, 'side': 'more', 'date': '2025-01-20'}])
        report = analyze_board(self.store, board['id'], as_of='2025-01-10T17:00:00Z')
        self.assertEqual(report['legs'][0]['n'], 9)
        self.assertEqual(report['legs'][0]['sample_last_date'], '2025-01-09')
        self.assertEqual(report['freshness']['fresh_feeds'], [])

    def test_research_lineups_use_exact_lines_distinct_players_and_teams(self):
        for player in ['Jalen Brunson', 'Josh Hart', 'Jayson Tatum', 'Derrick White']:
            for day in range(1, 13):
                self.observation(day, player=player)
        original = [
            {'player': 'Jalen Brunson', 'team': 'NYK', 'stat': 'pts', 'line': 5.5, 'side': 'more', 'date': '2025-01-20'},
            {'player': 'Jalen Brunson', 'team': 'NYK', 'stat': 'ast', 'line': 3.5, 'side': 'less', 'date': '2025-01-20'},
            {'player': 'Josh Hart', 'team': 'NYK', 'stat': 'reb', 'line': 4.5, 'side': 'less', 'date': '2025-01-20'},
            {'player': 'Jayson Tatum', 'team': 'BOS', 'stat': 'pts', 'line': 8.5, 'side': 'less', 'date': '2025-01-20'},
            {'player': 'Derrick White', 'team': 'BOS', 'stat': 'ast', 'line': 3.5, 'side': 'less', 'date': '2025-01-20'},
        ]
        board = self.store.create_board(original)
        report = analyze_board(self.store, board['id'], as_of='2025-01-19T17:00:00Z')
        self.assertEqual(len(report['suggested_entries']), 3)
        for suggestion in report['suggested_entries']:
            self.assertEqual(len({canonical_player(leg['player']) for leg in suggestion['legs']}), 2)
            self.assertEqual({leg['team'] for leg in suggestion['legs']}, {'NYK', 'BOS'})
            self.assertTrue(all(leg in board['legs'] for leg in suggestion['legs']))
            self.assertTrue(suggestion['research_only'])
            self.assertFalse(suggestion['actionable'])
            self.assertIsNone(suggestion['combined_win_probability'])
            self.assertIn('correlated', suggestion['correlation_warning'])
        self.assertEqual(self.store.get_board(board['id']), board)
        self.assertEqual(self.store.list_entries(), [])

    def test_all_knicks_or_missing_team_lines_require_other_team_metadata(self):
        for player in ['Jalen Brunson', 'Josh Hart']:
            for day in range(1, 13):
                self.observation(day, player=player)
        legs = [{'player': player, 'team': 'NYK', 'stat': 'pts', 'line': 8.5,
                 'side': 'more', 'date': '2025-01-20'} for player in ['Jalen Brunson', 'Josh Hart']]
        board = self.store.create_board(legs)
        report = analyze_board(self.store, board['id'], as_of='2025-01-19T17:00:00Z')
        self.assertEqual(report['suggested_entries'], [])
        self.assertIn('another team', report['lineup_note'])
        legs[0].pop('team')
        board = self.store.create_board(legs)
        report = analyze_board(self.store, board['id'], as_of='2025-01-19T17:00:00Z')
        self.assertEqual(report['suggested_entries'], [])
        self.assertIn('explicit team', report['lineup_note'])


if __name__ == '__main__':
    unittest.main()

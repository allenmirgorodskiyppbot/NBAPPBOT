"""Reproducible historical workbook views over canonical games and dated references.

This module deliberately never promotes season tracking means, source matchup
summaries or partial historical logs into canonical individual game records.
"""
from __future__ import annotations

import csv
import json
import os
import re
import shutil
import uuid
import warnings
from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn, TableStyleInfo
from openpyxl.worksheet.filters import AutoFilter

from .core import ET, digest, dumps, now

SUPPORTING_SHEETS = [
    "Data_PlayerSeason", "Data_PlayerGames", "Data_TeamGames", "Data_GameContext",
    "Data_TrackingGames", "Data_TrackingSeason", "Data_LineupsSeason",
    "Data_RotationStints", "Data_DefenderPairs", "Data_Shots", "Sources",
    "Data_Dictionary", "Data_Checks", "Missing_Data", "Update_Log", "Contents",
]
NBA_TEAMS = ("ATL BOS BKN CHA CHI CLE DAL DEN DET GSW HOU IND LAC LAL MEM MIA "
             "MIL MIN NOP NYK OKC ORL PHI PHX POR SAC SAS TOR UTA WAS").split()
STATS = ("minutes", "pts", "reb", "ast", "fg3m", "fg3a", "fgm", "fga", "ftm", "fta",
         "oreb", "dreb", "stl", "blk", "tov", "pf")
GAME_FIELDS = ["game_id", "player_id", "name", "team", "opponent", "season", "season_type", "date_et",
               "appeared", "started", "started_source", "started_known_at", "started_validation_status", "started_source_flags", "observation_scope", "minutes", "min_seconds", "source_minutes",
               "minutes_exact_text", "min_seconds_exact_text", "source_minutes_exact_text",
               "source_minute_precision", "source_minute_clock", "career_reference_source", "career_reference_known_at", "career_reference_minute_clock", "pts", "reb", "ast", "fg3m", "fg3a", "fgm", "fga",
               "ftm", "fta", "oreb", "dreb", "stl", "blk", "tov", "pf", "plus_minus",
               "known_at", "source", "hash"]
AUDIT_SHEETS = {
    "Source_catalog", "Missing_data_priorities", "Opponent_coverage", "Team_games_base_2025",
    "Team_games_advanced_2025", "Team_season_advanced_2025", "Passing_season_means_2025",
    "Rebound_season_means_2025", "Touches_season_means_2025", "Passing_one_date_2025",
    "Lineup_season_combinations", "Lineup_minute_checks", "Tracking_team_label_checks",
    "Sample_game_tracking", "Sample_game_box", "Sample_defender_pairs", "Sample_shot_events",
    "Unavailable_pair_history", "Older_Davis_vs_NYK",
}


def _read_reference(path):
    book = load_workbook(path, read_only=True, data_only=False)
    try:
        tables = {}
        for sheet in book:
            if sheet.title not in AUDIT_SHEETS:
                continue
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, ())
            if len(set(header)) != len(header) or any(not isinstance(h, str) for h in header):
                raise ValueError("Audit reference has invalid headers: " + sheet.title)
            rows = []
            for values in iterator:
                if not any(v is not None for v in values):
                    continue
                row = dict(zip(header, values))
                # NBA IDs in upstream references sometimes arrived as integers.
                # Game IDs are already text; never roundtrip those through int.
                for key, value in row.items():
                    if value is not None and (key.upper().endswith("_ID") or key in ("player_id", "game_id")):
                        row[key] = str(value)
                rows.append(row)
            tables[sheet.title] = rows
        return tables
    finally:
        book.close()


def archive_audit(store, path):
    """Copy an audited reference intact into durable data storage, by checksum."""
    path = Path(path)
    if path.is_dir():
        path = path / "Historical_Data_Audit.xlsx"
    tables = _read_reference(path)
    if not {"Source_catalog", "Opponent_coverage"}.issubset(tables):
        raise ValueError("Expected Historical_Data_Audit.xlsx with sources and opponent coverage")
    checksum = digest(path.read_bytes())
    root = store.root / "reference"
    root.mkdir(exist_ok=True)
    target = root / ("Historical_Data_Audit_" + checksum + ".xlsx")
    if not target.exists():
        temp = root / (".audit-" + uuid.uuid4().hex)
        shutil.copyfile(path, temp)
        if digest(temp.read_bytes()) != checksum:
            temp.unlink(missing_ok=True)
            raise ValueError("Audit archive checksum mismatch")
        os.replace(temp, target)
    elif digest(target.read_bytes()) != checksum:
        raise ValueError("Archived audit checksum mismatch")
    manifest = {"file": target.name, "sha256": checksum, "archived_at": now(),
                "tables": {k: len(v) for k, v in tables.items()},
                "meaning": "Dated historical reference; archive time does not refresh source observations"}
    temp = root / (".latest-" + uuid.uuid4().hex + ".json")
    temp.write_text(dumps(manifest))
    os.replace(temp, root / "latest.json")
    return manifest


def resolve_players(store, selected=None):
    selected = store.config.get("organized_players") if selected is None else selected
    if not isinstance(selected, list) or len(selected) not in (7, 8):
        raise ValueError("organized_players must name exactly seven or eight selected Knicks players (NBA IDs or exact names)")
    players = store.rows("SELECT id,name FROM players")
    result = []
    for supplied in selected:
        matches = [p for p in players if p["id"] == str(supplied) or p["name"].casefold() == str(supplied).casefold()]
        if len(matches) != 1:
            raise ValueError("Selected player is unknown or ambiguous: " + str(supplied))
        player = matches[0]
        # A historical Knicks affiliation is sufficient: this is the user's fixed
        # historical cohort, not a claim about current roster membership.
        if not store.one("SELECT 1 FROM player_games WHERE player_id=? AND team='NYK'", (player["id"],)):
            raise ValueError("Selected player has no verified Knicks history: " + player["name"])
        result.append(player)
    if len({p["id"] for p in result}) != len(selected):
        raise ValueError("The selected players must be unique")
    return result


resolve_selected_players = resolve_players


def _raw(store, sheet):
    # Repeated source imports must not multiply presentation summaries.
    unique = {}
    for row in store.rows("SELECT payload FROM raw_rows WHERE sheet=? ORDER BY file_hash,row_number", (sheet,)):
        payload = json.loads(row["payload"])
        unique[digest(payload)] = payload
    return list(unique.values())


def _safe_value(value):
    if isinstance(value, (dict, list, tuple)):
        return dumps(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return value


class _Writer:
    """Streaming workbook with independently filterable, nonoverlapping tables."""
    def __init__(self, csv_dir=None):
        self.book = Workbook(write_only=True)
        self.positions = {}
        self.tables = []
        self.contents = []
        self.counter = 0
        self.csv_dir = Path(csv_dir) if csv_dir else None
        if self.csv_dir:
            self.csv_dir.mkdir(parents=True, exist_ok=True)

    def sheet(self, name, purpose):
        ws = self.book.create_sheet(name)
        # A write-only sheet serializes its pane configuration at first append.
        # The first table header is on row5 or6, so freeze through row6 now.
        ws.freeze_panes = "D7"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.column_dimensions["A"].width = 22
        for i in range(2, 85):
            ws.column_dimensions[get_column_letter(i)].width = 18
        self.positions[name] = 0
        self.contents.append({"sheet": name, "purpose": purpose, "link": name + "!A1"})
        self.line(ws, [purpose], bold=True)
        self.line(ws, ["Return to Contents"], hyperlink="#'Contents'!A1")
        return ws

    def line(self, ws, values, bold=False, hyperlink=None):
        cells = []
        for value in values:
            cell = WriteOnlyCell(ws, value=_safe_value(value))
            if isinstance(cell.value, str):
                cell.data_type = "s"
            if bold:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="1F4E78")
            if hyperlink:
                cell.hyperlink = hyperlink
                cell.font = Font(color="0563C1", underline="single")
            cells.append(cell)
        ws.append(cells)
        self.positions[ws.title] += 1

    def table(self, ws, title, rows, headers=None, note=None):
        rows = list(rows)
        headers = list(headers or dict.fromkeys(k for r in rows for k in r)) or ["availability"]
        self.line(ws, [])
        self.line(ws, [title], bold=True)
        if note:
            self.line(ws, [note])
        start = self.positions[ws.title] + 1
        self.line(ws, headers, bold=True)
        if rows:
            for row in rows:
                self.line(ws, [row.get(h) for h in headers])
        else:
            self.line(ws, ["Unavailable; not zero"])
        self.counter += 1
        name = "T_" + str(self.counter)
        ref = "A" + str(start) + ":" + get_column_letter(len(headers)) + str(self.positions[ws.title])
        table = Table(displayName=name, ref=ref,
                      tableColumns=[TableColumn(id=i + 1, name=h) for i, h in enumerate(headers)])
        table.autoFilter = AutoFilter(ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        # write-only sheets need explicit column names; this avoids automatic
        # header discovery and keeps numeric-looking identifiers stored as text.
        with warnings.catch_warnings():
            warnings.filterwarnings("ignore", message="In write-only mode you must add table columns manually")
            ws.add_table(table)
        self.tables.append({"sheet": ws.title, "table": title, "range": ref, "rows": len(rows), "columns": len(headers)})
        if self.csv_dir:
            filename = re.sub(r"[^A-Za-z0-9_-]+", "_", ws.title + "__" + title)[:180] + ".csv"
            with (self.csv_dir / filename).open("w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow(headers)
                for row in rows:
                    values = [_safe_value(row.get(h)) for h in headers]
                    writer.writerow(["'" + v if isinstance(v, str) and v.startswith(("=", "+", "-", "@", "\t", "\r")) else v for v in values])


def _summary(rows, **labels):
    played = [r for r in rows if r.get("appeared") == 1]
    dates = [r["date_et"] for r in played]
    result = {**labels, "played_games": len(played), "listed_nonappearances": sum(r.get("appeared") == 0 for r in rows),
              "first_game": min(dates) if dates else None, "through_date": max(dates) if dates else None,
              "actual_teams": ", ".join(sorted({r["team"] for r in played if r.get("team")})),
              "basis": "Verified unique played games in the stated phase; totals require complete observations for the field"}
    for stat in STATS:
        observations = [r[stat] for r in played if r.get(stat) is not None]
        total = sum(observations) if observations else None
        result[stat + "_observed"] = len(observations)
        result[stat + "_total"] = total if len(observations) == len(played) and played else None
        result[stat + "_mean_observed"] = total / len(observations) if observations else None
    for made, attempted, label in (("fgm", "fga", "fg_pct"), ("fg3m", "fg3a", "fg3_pct"), ("ftm", "fta", "ft_pct")):
        m, a = result[made + "_total"], result[attempted + "_total"]
        result[label] = m / a if m is not None and a is not None and a else None
    minutes = result["minutes_total"]
    for stat in ("pts", "reb", "ast", "fg3m", "fg3a", "fga", "fta"):
        total = result[stat + "_total"]
        result[stat + "_per36"] = total / minutes * 36 if total is not None and minutes else None
    return result


def _group_summaries(rows, keys, **labels):
    grouped = defaultdict(list)
    for row in rows:
        grouped[tuple(row.get(k) for k in keys)].append(row)
    summaries = [_summary(value, **labels, **dict(zip(keys, key))) for key, value in grouped.items()]
    return sorted(summaries, key=lambda r: (str(r.get("season", "")), str(r.get("name", "")), str(r.get("team", ""))), reverse=True)


def _game_rows(store):
    rows = store.rows("SELECT pg.*,p.name FROM player_games pg JOIN players p ON pg.player_id=p.id ORDER BY date_et,game_id,player_id")
    for row in rows:
        payload = json.loads(row.pop("payload") or "{}")
        row["source_minutes"] = next((payload[k] for k in ("Original reported MIN", "MIN_SEC", "Raw source MIN_SEC", "Minutes (source precision)", "MIN") if payload.get(k) is not None), None)
        # XLSX numeric writers/Excel can round long binary float displays. Keep a
        # text companion containing the full canonical repr for exact recovery.
        for field in ("minutes", "min_seconds", "source_minutes"):
            row[field + "_exact_text"] = str(row[field]) if row.get(field) is not None else None
        row["source_minute_precision"] = payload.get("Minute precision") or "As supplied by source; canonical precision retained"
    return rows


def _context(store, games):
    contexts = []
    completed_by_team = defaultdict(list)
    rest = {}
    for game in sorted(games, key=lambda r: (r["date_et"], r["id"])):
        context = {k: v for k, v in game.items() if k != "metadata"}
        context["game_id"] = context.pop("id")
        context["tipoff_et"] = datetime.fromisoformat(game["tipoff"]).astimezone(ET).isoformat() if game.get("tipoff") else None
        for side in ("home", "away"):
            team = game.get(side)
            previous = completed_by_team.get(team, [])
            prior = previous[-1] if previous else None
            days = (datetime.fromisoformat(game["date_et"]) - datetime.fromisoformat(prior["date_et"])).days - 1 if prior else None
            context[side + "_previous_game_id"] = prior["id"] if prior else None
            context[side + "_rest_days"] = max(days, 0) if days is not None else None
            context[side + "_back_to_back"] = days == 0 if days is not None else None
            rest[(game["id"], team)] = context[side + "_rest_days"]
        context["rest_coverage"] = "Derived from available canonical team schedule; first known game has unknown prior rest"
        contexts.append(context)
        if game.get("status") == "completed":
            for team in (game.get("home"), game.get("away")):
                if team:
                    completed_by_team[team].append(game)
    return contexts, rest


def _reference(store, audit_dir=None):
    if audit_dir:
        path = Path(audit_dir)
        if path.is_dir():
            path = path / "Historical_Data_Audit.xlsx"
        return _read_reference(path), {"file": str(path), "sha256": digest(path.read_bytes())}
    manifest = store.root / "reference" / "latest.json"
    if not manifest.exists():
        return {}, None
    metadata = json.loads(manifest.read_text())
    path = manifest.parent / metadata["file"]
    if path.name != metadata["file"] or digest(path.read_bytes()) != metadata["sha256"]:
        raise ValueError("Historical reference archive checksum mismatch")
    return _read_reference(path), metadata


def _annotate(rows, dataset, sources, units, season="2025-26"):
    available = [r for r in sources if r.get("dataset") == dataset]
    source = available[0] if available else {}
    result = []
    for row in rows:
        item = {"reference_dataset": dataset, "coverage_season": season, "units": units, **row}
        if dataset == "season_five_player_lineups":
            per_team = [r for r in sources if r.get("dataset") == "season_lineups_" + str(row.get("TEAM_ABBREVIATION", ""))]
            if per_team:
                source = per_team[0]
        item["source_url"] = source.get("source_url")
        item["retrieved_at"] = source.get("retrieved_at")
        result.append(item)
    return result


def _older_annual(records):
    """Summarize only fields actually present in separate older source logs."""
    groups = defaultdict(list)
    for row in records:
        groups[str(row.get("Season", "Unknown"))].append(row)
    result = []
    for season in sorted(groups, reverse=True):
        group = groups[season]
        dates = [str(r.get("Game date"))[:10] for r in group if r.get("Game date")]
        summary = {"season": season, "listed_source_games": len(group),
                   "first_game": min(dates) if dates else None, "through_date": max(dates) if dates else None,
                   "actual_teams": ", ".join(sorted({str(r["Actual team"]) for r in group if r.get("Actual team")})),
                   "scope": "Separate uploaded historical logs; source coverage claims retained, canonical game IDs/minutes unavailable"}
        for source, target in (("Points", "pts"), ("Rebounds", "reb"), ("Assists", "ast"), ("3PM", "fg3m"), ("3PA", "fg3a")):
            observed = [r[source] for r in group if isinstance(r.get(source), (int, float))]
            summary[target + "_observed"] = len(observed)
            summary[target + "_total"] = sum(observed) if len(observed) == len(group) else None
            summary[target + "_mean_observed"] = sum(observed) / len(observed) if observed else None
        summary["minutes"] = None
        summary["minutes_availability"] = "Not supplied in these archival rows; not zero"
        result.append(summary)
    return result


def _warning_rows(audit, selected_names):
    rows = [
        {"code": "career_coverage", "availability": "Covered seasons only", "detail": "Available canonical seasons are not necessarily an entire career; retain each source's coverage labels."},
        {"code": "tracking_backfill", "availability": "Collected observations only", "detail": "Season means retain their stated season and cutoff. Only collected and verified game observations fill game-level coverage."},
        {"code": "rotation_stints", "availability": "Unavailable scopes remain missing", "detail": "Failed or unavailable game rotations cannot be filled from season lineup averages."},
        {"code": "starting_roles", "availability": "Partial", "detail": "Preserve known starting indicators and sample box-score starting positions; missing starts remain unknown."},
        {"code": "defender_rebounds", "availability": "Unavailable; not zero", "detail": "Measured defender-pair source does not contain defender-specific rebounds."},
        {"code": "shot_defender_distance", "availability": "Unavailable; not zero", "detail": "Sample shot chart has no measured per-shot defender distance."},
        {"code": "historical_asof", "availability": "Insufficient historical pregame vintages", "detail": "Retrieval today is not evidence a feature was known before a historical game. Final season means are not pregame features."},
        {"code": "estimated_assignments", "availability": "Estimates require supplied rules", "detail": "Primary/alternate defender rules, when supplied, remain estimates. A current roster does not establish a measured defensive assignment."},
    ]
    if audit:
        lineups = audit.get('Lineup_season_combinations', [])
        if lineups:
            rows.append({'code': 'lineup_cap', 'availability': 'Archived per-team requests',
                'detail': f"The supplied audit contains {len(lineups)} per-team season combinations. Per-team requests avoid truncating a global request; combinations are not individual absence splits or rotation stints."})
        tracking_labels = audit.get('Tracking_team_label_checks', [])
        if tracking_labels:
            rows.append({'code': 'tracking_team_labels', 'availability': 'Source discrepancies retained',
                'detail': f"The supplied audit records {len(tracking_labels)} tracking-team label disagreements. Raw labels and verified historical joins remain separate; see the archived checks."})
        if audit.get('Lineup_minute_checks'):
            rows.append({'code': 'lineup_minute_gaps', 'availability': 'Source checks retained',
                'detail': 'Supplied lineup minute checks retain their source differences and precision. No source minutes are changed to force a pass.'})
    else:
        rows.append({"code": "audit_reference", "availability": "Unavailable", "detail": "The private uploaded audit is not included in this data package. Import it privately to add its historical reference tables; no archived counts or completed source checks are assumed."})
    return rows


def _supplemental_rows(table):
    """Keep refresh metadata attached without changing raw NBA field units."""
    metadata = {key: table.get(key) for key in ("key", "season", "scope", "unit", "source_url", "retrieved_at",
        "coverage_cutoff", "coverage_note", "publisher_updated_at", "refresh_status", "warnings")}
    return [{"reference_dataset": metadata["key"], **metadata, **row} for row in table.get("rows", [])]


def _load_references(store):
    import importlib
    defaults = {
        'career': ('career_reference', 'load_career_reference', {'rows': [], 'tables': [], 'controls': [], 'errors': [], 'status': 'not_collected'}),
        'opportunities': ('opportunity_reference', 'load_opportunity_reference', {'tracking_games': [], 'passing_games': [], 'starts': [], 'sources': [], 'checks': [], 'errors': [], 'coverage': {}}),
        'rotations': ('rotation_reference', 'load_rotation_reference', {'stints': [], 'lineup_intervals': [], 'starts': [], 'sources': [], 'checks': [], 'errors': [], 'coverage': {}}),
        'matchups': ('matchup_reference', 'load_matchup_reference', {'matchups': [], 'shots': [], 'sources': [], 'coverage': [], 'errors': [], 'summary': {}}),
    }
    result = {}
    for kind, (module, function, empty) in defaults.items():
        try:
            imported = importlib.import_module('.' + module, __package__)
        except ModuleNotFoundError as exc:
            if exc.name != __package__ + '.' + module:
                raise
            result[kind] = empty
        else:
            result[kind] = getattr(imported, function)(store)
    return result


def _phase(row):
    return row.get('season_type') or row.get('phase') or ('Playoffs' if str(row.get('game_id', '')).startswith('004') else 'Regular Season')


def _prepare_game(row, scope):
    row = dict(row)
    row['game_id'], row['player_id'] = str(row['game_id']), str(row['player_id'])
    row['season_type'] = _phase(row)
    row['observation_scope'] = scope
    for field in STATS:
        row.setdefault(field, None)
    row.setdefault('started', None)
    row.setdefault('appeared', 1)
    for field in ('minutes', 'min_seconds', 'source_minutes'):
        row.setdefault(field + '_exact_text', str(row[field]) if row.get(field) is not None else None)
    return row


def _merge_scoped_games(canonical, career, scope, starts):
    selected = set(scope['selected_ids'])
    by_team = {team: set(ids) for team, ids in scope['by_opponent'].items()}
    result, conflicts = {}, []
    for source in career:
        if str(source.get('player_id')) not in selected:
            continue
        row = _prepare_game(source, 'Verified selected-player career reference')
        key = row['game_id'], row['player_id']
        if key in result and result[key] != row:
            raise ValueError('Duplicate career reference observations disagree: ' + str(key))
        result[key] = row
    for source in canonical:
        pid = source['player_id']
        if pid not in selected and not (source.get('opponent') == 'NYK' and pid in by_team.get(source.get('team'), set())):
            continue
        row = _prepare_game(source, 'Canonical database; selected player or exact named matchup vs NYK')
        key = row['game_id'], pid
        other = result.get(key)
        if other:
            differences = {k: {'canonical': row.get(k), 'career_reference': other.get(k)} for k in STATS if k != 'minutes' and row.get(k) is not None and other.get(k) is not None and row[k] != other[k]}
            if row['season_type'] != other['season_type']:
                raise ValueError('Career and canonical game phases disagree')
            if differences:
                conflicts.append({'game_id': key[0], 'player_id': pid, 'action': 'Canonical source retained', 'differences': differences})
            row['career_reference_source'] = other.get('source')
            row['career_reference_known_at'] = other.get('known_at')
            row['career_reference_minute_clock'] = other.get('source_minute_clock')
        result[key] = row
    for source in starts:
        key = str(source.get('game_id')), str(source.get('player_id'))
        status = str(source.get('validation_status', '')).lower()
        row = result.get(key)
        # Traditional box rows can supply named opponents in playoff fixtures
        # absent from the regular-season canonical table. Validate their actual
        # full-game fields before adding a separate reference observation.
        if row is None and key[1] not in selected and source.get('opponent') == 'NYK' and key[1] in by_team.get(source.get('team'), set()):
            if source.get('appeared') not in (0, 1) or not all(field in source for field in ('pts', 'reb', 'ast', 'fgm', 'fga', 'fg3m', 'fg3a', 'ftm', 'fta')):
                continue
            if any(word in status for word in ('conflict', 'failed', 'unverified')):
                continue
            from .ingest import normalize_game
            incoming = {**source, 'Player': source.get('name'), 'source': source.get('source_url'), 'known_at': source.get('retrieved_at')}
            normalized = normalize_game(incoming, source.get('retrieved_at'))
            normalized.pop('payload', None)
            row = _prepare_game({**source, **normalized, 'name': source.get('name'), 'source_minutes': source.get('source_minutes'), 'season_type': _phase(source)}, 'Verified named-opponent full-game box reference against NYK')
            if row.get('appeared') == 1 and row.get('started') in (0, 1):
                row.update(started_source=source.get('source_url'), started_known_at=source.get('retrieved_at'),
                           started_validation_status=source.get('role_validation_status') or source.get('validation_status'),
                           started_source_flags=source.get('flags', []))
            result[key] = row
        if row is None or row.get('appeared') != 1 or source.get('started') not in (0, 1):
            continue
        role_verified = source.get('role_validation_status') == 'verified_official_game_identity_and_five_starters_per_team'
        appearance_conflict = 'source_appearance_difference' in source.get('flags', [])
        if appearance_conflict or (not role_verified and any(word in status for word in ('conflict', 'failed', 'unverified'))):
            continue
        if row.get('started') is None:
            row['started'] = int(source['started'])
            row['started_source'] = source.get('source_url') or source.get('box_source_url')
            row['started_known_at'] = source.get('retrieved_at') or source.get('box_retrieved_at')
            row['started_validation_status'] = source.get('role_validation_status') or source.get('validation_status') or 'Validated rotation and opening box lineup'
            row['started_source_flags'] = source.get('flags', [])
        elif row['started'] != source['started']:
            conflicts.append({'game_id': key[0], 'player_id': key[1], 'action': 'Existing role retained; source disagreement flagged', 'existing_started': row['started'], 'new_started': source['started'], 'source': source.get('source_url')})
    return sorted(result.values(), key=lambda r: (r['date_et'], r['game_id'], r['player_id'])), conflicts


def _scoped_correction_audit(records, game_rows):
    """Show only corrections for exported observations or their actual games.

    An opponent's annual control or another fixture is not a matchup observation.
    The underlying audit database and preserved source archive remain complete.
    """
    keys = {(str(row['game_id']), str(row['player_id'])) for row in game_rows}
    game_ids = {key[0] for key in keys}
    output = []
    for record in records:
        parts = str(record.get('record_key') or '').split('/')
        if (len(parts) == 2 and tuple(parts) in keys) or (len(parts) == 1 and parts[0] in game_ids):
            output.append(record)
    return output


def _pair_rows(pairs, game_rows, baseline_rows, measured):
    result = []
    for pair in pairs:
        kid, oid, team = pair['knicks_id'], pair['opponent_id'], pair['opponent_team']
        measurements = [r for r in measured if r.get('selected_knicks_player_id') == kid and r.get('matchup_player_id') == oid and r.get('opponent_team') == team]
        for phase in ('Regular Season', 'Playoffs'):
            k_history = [r for r in game_rows if r['player_id'] == kid and r.get('team') == 'NYK' and r.get('opponent') == team and _phase(r) == phase]
            o_history = [r for r in game_rows if r['player_id'] == oid and r.get('team') == team and r.get('opponent') == 'NYK' and _phase(r) == phase]
            tracked = [r for r in measurements if _phase(r) == phase]
            if phase == 'Playoffs' and not (k_history or o_history or tracked):
                continue
            ks, os = _summary(k_history), _summary(o_history)
            row = {**pair, 'season_type': phase,
                   'whole_game_scope': 'Separate full-game production against this team; not production while guarded by the named player',
                   'knicks_full_game_samples': ks['played_games'], 'opponent_full_game_samples': os['played_games'],
                   'measured_pair_games': len({r['game_id'] for r in tracked}),
                   'pair_scope': 'Measured partial-possession counts are separate from full-game samples; estimated future assignment remains unconfirmed'}
            for prefix, summary in (('knicks_full_game_', ks), ('opponent_full_game_', os)):
                for stat in ('minutes', 'pts', 'reb', 'ast', 'fg3m', 'fg3a', 'fga', 'fta'):
                    row[prefix + stat + '_mean'] = summary[stat + '_mean_observed']
            for prefix, pid in (('knicks_baseline_', kid), ('opponent_baseline_', oid)):
                available = [r for r in baseline_rows if r['player_id'] == pid and _phase(r) == phase]
                season = max((r['season'] for r in available), default=None)
                baseline = _summary([r for r in available if r['season'] == season])
                row[prefix + 'season'] = season
                row[prefix + 'played_games'] = baseline['played_games']
                row[prefix + 'actual_teams'] = baseline['actual_teams']
                for stat in ('minutes', 'pts', 'reb', 'ast', 'fg3m', 'fg3a', 'fga', 'fta'):
                    row[prefix + stat + '_mean'] = baseline[stat + '_mean_observed']
            for direction in ('offense', 'defense'):
                group = [r for r in tracked if r.get('knicks_side') == direction]
                row[direction + '_measured_games'] = len({r['game_id'] for r in group})
                for field in ('partialPossessions', 'playerPoints', 'matchupFieldGoalsAttempted', 'matchupThreePointersAttempted', 'matchupPotentialAssists'):
                    values = [r.get(field) for r in group]
                    row[direction + '_' + field] = sum(values) if values and all(isinstance(v, (int, float)) for v in values) else None
            result.append(row)
    return result


def build_organized_workbook(store, destination, audit_dir=None, selected_players=None, csv_dir=None):
    """Build selected-player and exact named-matchup views without exporting rosters."""
    selected = resolve_players(store, selected_players)
    from .ingest import validate
    from .scope import matchup_scope, normalized_name
    canonical_validation = validate(store)
    if canonical_validation['status'] == 'failed':
        raise ValueError('Resolve canonical validation errors before generating the organized workbook')
    scope = matchup_scope(store, selected=selected)
    selected_ids, all_ids = set(scope['selected_ids']), set(scope['all_ids'])
    by_team = {team: set(ids) for team, ids in scope['by_opponent'].items()}
    selected_names = {p['name'] for p in selected}
    identities = store.rows('SELECT id,name FROM players')
    names = {p['id']: p['name'] for p in identities}
    named_names = {normalized_name(names[pid]) for pid in all_ids if pid in names}
    name_ids = {normalized_name(p['name']): p['id'] for p in identities}
    destination = Path(destination)
    if destination.suffix.lower() != '.xlsx':
        destination = destination / 'Knicks_Historical_Organized.xlsx'
    for original in store.rows('SELECT original_path FROM files'):
        if original['original_path'] and destination.resolve() == Path(original['original_path']).resolve():
            raise ValueError('Refusing to overwrite a preserved original workbook')
    destination.parent.mkdir(parents=True, exist_ok=True)
    audit, audit_meta = _reference(store, audit_dir)
    sources = list(audit.get('Source_catalog', []))
    from .historical_context import load_context_reference
    supplemental = load_context_reference(store)
    references = _load_references(store)
    career, opportunities, rotations, matchups = (references[k] for k in ('career', 'opportunities', 'rotations', 'matchups'))
    career_rows = career.get('rows')
    if career_rows is None:
        career_rows = [row for table in career.get('tables', []) for row in table.get('rows', [])]
    career_rows = [row for row in career_rows if str(row.get('player_id')) in selected_ids]
    career_controls = [row for row in career.get('controls', []) if str(row.get('player_id')) in selected_ids]
    canonical = _game_rows(store)
    canonical_by_key = {(r['game_id'], r['player_id']): r for r in canonical}
    rows, merge_conflicts = _merge_scoped_games(canonical, career_rows, scope, opportunities.get('starts', []) + rotations.get('starts', []))
    selected_regular = [r for r in rows if r['player_id'] in selected_ids and _phase(r) == 'Regular Season']
    selected_playoffs = [r for r in rows if r['player_id'] in selected_ids and _phase(r) == 'Playoffs']
    baseline_rows = selected_regular + selected_playoffs + [_prepare_game(r, 'Named matchup broader season baseline only') for r in canonical if r['player_id'] in set(scope['opponent_ids'])]
    annual = _group_summaries(baseline_rows, ['player_id', 'name', 'season', 'season_type'])
    regular_annual = [r for r in annual if r['season_type'] == 'Regular Season']
    by_player = defaultdict(list)
    for row in rows:
        by_player[row['player_id']].append(row)
    games = store.rows('SELECT * FROM games ORDER BY date_et,id')
    game_map = {g['id']: g for g in games}
    contexts, rest = _context(store, games)
    game_ids = {r['game_id'] for r in rows}
    contexts = [r for r in contexts if r['game_id'] in game_ids or r.get('home') == 'NYK' or r.get('away') == 'NYK']
    raw_stats = [r for r in _raw(store, 'Player Stats') if normalized_name(r.get('Player', '')) in named_names]
    raw_career = [r for r in _raw(store, 'Career Rest Games') if r.get('Player') in selected_names]
    raw_rest = [r for r in _raw(store, 'Rest History') if r.get('Player') in selected_names]
    assignments = [r for r in _raw(store, 'Matchups') if r.get('Knicks player') in selected_names]
    position_defense = _raw(store, 'Position Defense')
    pair_keys = {(p['knicks_id'], p['opponent_team'], p['opponent_id']) for p in scope['pairs']}

    def valid_pair(row):
        return (str(row.get('selected_knicks_player_id')), row.get('opponent_team'), str(row.get('matchup_player_id'))) in pair_keys

    measured = [r for r in matchups.get('matchups', []) if valid_pair(r)]
    def allowed_person(row):
        pid = str(row.get('player_id', row.get('PLAYER_ID', row.get('personId', ''))))
        if pid in selected_ids:
            return True
        game_id = str(row.get('game_id', row.get('GAME_ID', row.get('DERIVED_GAME_ID', ''))))
        existing = canonical_by_key.get((game_id, pid))
        if existing:
            return existing.get('opponent') == 'NYK' and pid in by_team.get(existing.get('team'), set())
        team = row.get('team') or row.get('ACTUAL_HISTORICAL_TEAM') or row.get('TEAM_ABBREVIATION')
        opponent = row.get('opponent') or row.get('opponent_team')
        return (opponent == 'NYK' or (row.get('side') == 'Named matchup opponent' and team == row.get('opponent_team'))) and pid in by_team.get(team, set())

    def raw_pair_allowed(row):
        offense, defense = name_ids.get(normalized_name(row.get('Offensive player', ''))), name_ids.get(normalized_name(row.get('Defender', '')))
        team = row.get('Opponent team')
        return (offense, team, defense) in pair_keys or (defense, team, offense) in pair_keys

    defenders = [r for r in _raw(store, 'Defender History') if raw_pair_allowed(r)]
    observed_rs_pairs = defaultdict(set)
    for measurement in measured:
        if _phase(measurement) == 'Regular Season':
            observed_rs_pairs[(str(measurement.get('offensive_player_id')), str(measurement.get('defender_player_id')), measurement.get('opponent_team'))].add(measurement['game_id'])
    unavailable_original = []
    for original_pair in audit.get('Unavailable_pair_history', []):
        if not raw_pair_allowed(original_pair):
            continue
        key = (name_ids.get(normalized_name(original_pair.get('Offensive player', ''))), name_ids.get(normalized_name(original_pair.get('Defender', ''))), original_pair.get('Opponent team'))
        games_observed = len(observed_rs_pairs.get(key, set()))
        unavailable_original.append({**original_pair, 'new_window_regular_season_pair_games': games_observed,
            'new_window_note': 'New measured games in the collected window; not a replacement for the originally unavailable career aggregate' if games_observed else 'No observed regular-season pair in the collected window; missing is not zero production'})
    recovered_original_pairs = sum(r['new_window_regular_season_pair_games'] > 0 for r in unavailable_original)
    tracking_games = [r for r in opportunities.get('tracking_games', []) if allowed_person(r)]
    passing_games = [r for r in opportunities.get('passing_games', []) if allowed_person(r)]
    observed_starts = [r for r in opportunities.get('starts', []) + rotations.get('starts', []) if allowed_person(r)]
    stints = [r for r in rotations.get('stints', []) if allowed_person(r)]
    lineup_intervals = [{**r, 'scoped_player_ids': [str(pid) for pid in r.get('scoped_player_ids', []) if str(pid) in all_ids]} for r in rotations.get('lineup_intervals', []) if set(map(str, r.get('scoped_player_ids', []))) & all_ids]
    shots = [r for r in matchups.get('shots', []) if allowed_person(r)]
    pair_rows = _pair_rows(scope['pairs'], rows, baseline_rows, measured)
    supplemental_tables = supplemental.get('tables', [])
    team_tables = [t for t in supplemental_tables if t.get('kind') == 'team_games']
    tracking_tables = [t for t in supplemental_tables if t.get('kind') == 'tracking_season']
    team_scopes = {(t.get('season'), t.get('measure')) for t in team_tables}
    tracking_scopes = {(t.get('season'), t.get('measure')) for t in tracking_tables}
    opportunity_rows = []
    for sheet, dataset, measure in (('Passing_season_means_2025', 'season_passing', 'Passing'), ('Rebound_season_means_2025', 'season_rebounding', 'Rebounding'), ('Touches_season_means_2025', 'season_touches', 'Possessions')):
        if ('2025-26', measure) not in tracking_scopes:
            opportunity_rows.extend(_annotate([r for r in audit.get(sheet, []) if str(r.get('PLAYER_ID')) in all_ids], dataset, sources, 'Published per-game season means; not game observations'))
    for table in tracking_tables:
        opportunity_rows.extend(r for r in _supplemental_rows(table) if str(r.get('PLAYER_ID')) in all_ids)
    for table in supplemental_tables + career.get('tables', []):
        sources.append({k: v for k, v in table.items() if k not in ('rows', 'raw_source_response', 'checks')} | {'row_count': len(table.get('rows', []))})
    for reference in (opportunities, rotations, matchups):
        sources.extend(reference.get('sources', []))
    warnings = _warning_rows(audit, selected_names)
    warnings = [w for w in warnings if w['code'] not in ('career_coverage', 'tracking_backfill', 'rotation_stints', 'starting_roles', 'unavailable_defender_pairs')]
    played = [r for r in rows if r['player_id'] in selected_ids and r['appeared'] == 1]
    unknown_starts = sum(r.get('started') is None for r in played)
    canonical_selected_played = [r for r in canonical if r['player_id'] in selected_ids and r['appeared'] == 1]
    original_unknown_starts = sum(r.get('started') is None for r in canonical_selected_played)
    merged_by_key = {(r['game_id'], r['player_id']): r for r in rows}
    remaining_canonical_unknown_starts = sum(merged_by_key[(r['game_id'], r['player_id'])].get('started') is None for r in canonical_selected_played)
    warnings += [
        {'code': 'unavailable_defender_pairs', 'availability': 'Original source gaps retained; new window checked' if audit else 'Archived pair inventory unavailable', 'detail': f"{len(unavailable_original)} original regular-season pair aggregates remain dated unavailable claims. {recovered_original_pairs} of these directed pairs have new measured games within the collected regular-season window; {len(unavailable_original) - recovered_original_pairs} have no observation in that window. New samples are not complete career replacements, and absent observations are not zero." if audit else 'No private archived pair inventory was supplied. Any newly measured pairs have their own coverage; an absent inventory is not a verified zero-gap result.'},
        {'code': 'career_coverage', 'availability': career.get('status', 'not_collected'), 'detail': f"{len(career_rows)} verified reference game rows supplement canonical observations with canonical precedence. Regular season and playoffs remain separate; quarantined packets and older unmatched archives remain flagged."},
        {'code': 'tracking_backfill', 'availability': 'Scoped game data; see source coverage', 'detail': f"{len(tracking_games)} game tracking rows and {len(passing_games)} verified single-game passing rows are included for selected players/named matchups. Season averages stay separate; endpoint/game failures are not zeros."},
        {'code': 'starting_roles', 'availability': 'Partial' if unknown_starts else 'Verified within exported played games', 'detail': f"{unknown_starts} of {len(played)} selected-player played games still lack a verified start/bench indicator. Season GS totals are not game-level roles; new role retrieval timestamps remain separate from original statistic knowledge times."},
        {'code': 'rotation_stints', 'availability': rotations.get('status', 'not_collected'), 'detail': f"{len(stints)} scoped validated player stints and {len(lineup_intervals)} five-player team context intervals are included. Only validated games are published; team context does not establish individual absence effects."},
        {'code': 'named_matchup_scope', 'availability': 'Explicit scope' if scope['pairs'] else 'Unavailable', 'detail': f"{len(scope['pairs'])} named primary/alternate assignments define exact pairs. Opponent game logs are restricted to actual games against NYK; no complete opponent rosters are exported." if scope['pairs'] else 'No primary/alternate matchup rules have been imported. A fresh deployment does not include the private uploaded workbook or its custom matchup assignments. Named pair views remain unavailable until those rules are imported privately.'},
    ]
    errors = []
    for kind, reference in {'context': supplemental, **references}.items():
        for error in reference.get('errors', []):
            errors.append({'reference_kind': kind, **error})
            warnings.append({'code': kind + '_source_failure', 'availability': 'Unavailable or retained dated source', 'detail': dumps(error)})
    for item in scope['unresolved']:
        warnings.append({'code': 'unresolved_matchup_identity', 'availability': 'Not joined', 'detail': dumps(item)})
    writer = _Writer(csv_dir)
    created = now()
    last_success = store.one("SELECT finished_at FROM update_runs WHERE status='success' ORDER BY finished_at DESC LIMIT 1")
    cutoff = max((r['date_et'] for r in canonical), default=None)
    coverage_note = f"Generated {created}; canonical completed-game coverage through {cutoff}. Newly retrieved historical sources are retrospective, not pregame publication vintages."
    for player in selected:
        pid = player['id']
        surname = re.sub(r'[^A-Za-z0-9_]', '', player['name'].split()[-1])
        sheet_name = ('NYK_' + surname)[:31]
        if sheet_name in writer.positions:
            sheet_name = ('NYK_' + surname[:17] + '_' + pid)[:31]
        ws = writer.sheet(sheet_name, player['name'] + ' — history and exact named matchup pairs')
        writer.table(ws, 'Identity and coverage', [{**player, 'current_season': store.config['season'], 'generated_at': created, 'canonical_through': cutoff, 'last_successful_refresh': last_success['finished_at'] if last_success else None, 'coverage_note': coverage_note}])
        history = by_player[pid]
        regular = [r for r in history if _phase(r) == 'Regular Season']
        playoffs = [r for r in history if _phase(r) == 'Playoffs']
        writer.table(ws, 'Annual regular-season summaries — current season first', [r for r in regular_annual if r['player_id'] == pid], note='Canonical and verified career games are deduplicated by game ID/player ID. Source-minute precision remains visible; actual traded teams are preserved.')
        writer.table(ws, 'Annual playoff summaries — separate phase', [r for r in annual if r['player_id'] == pid and r['season_type'] == 'Playoffs'])
        covered_dates = {r['date_et'] for r in regular}
        older_history = [r for r in raw_career if r.get('Player') == player['name'] and str(r.get('Game date', ''))[:10] not in covered_dates]
        writer.table(ws, 'Unmatched older annual archives — separate coverage', _older_annual(older_history), note='Only archival dates without a verified regular-season game remain. These observations lack canonical game IDs and are not added to verified totals.')
        writer.table(ws, 'Current regular-season totals and averages', [_summary([r for r in regular if r['season'] == store.config['season']], player_id=pid, name=player['name'], season=store.config['season'], season_type='Regular Season')])
        writer.table(ws, 'Covered career regular-season totals', [_summary(regular, player_id=pid, name=player['name'], coverage='Verified covered regular-season games; see independent controls and source failures')])
        writer.table(ws, 'Covered career playoff totals — separate phase', [_summary(playoffs, player_id=pid, name=player['name'], coverage='Verified covered playoff games only')])
        writer.table(ws, 'Independent published annual controls', [r for r in career_controls if str(r.get('player_id')) == pid], note='Published GS is an annual total, not a game-level starting indicator. Team/TOT overlap has been reconciled by the source loader.')
        writer.table(ws, 'Uploaded career summary — dated source coverage', [r for r in raw_stats if r.get('Player') == player['name'] and 'career' in str(r.get('Period', '')).lower()])
        recent = sorted((r for r in regular if r['appeared'] == 1), key=lambda r: (r['date_et'], r['game_id']), reverse=True)
        writer.table(ws, 'Recent regular-season played games 5 10 20', [_summary(recent[:n], requested_window=n, player_id=pid) for n in (5, 10, 20)])
        writer.table(ws, 'Specific named matchup pairs — no roster cross product', [r for r in pair_rows if r['knicks_id'] == pid], note='Only this player’s primary and alternate names are paired. Whole-game H2H production and measured partial-possession guarding samples are different observations.')
        writer.table(ws, 'Measured game-level named pair history', [r for r in measured if r.get('selected_knicks_player_id') == pid])
        writer.table(ws, 'Available game-level opportunities', [r for r in tracking_games + passing_games if str(r.get('player_id')) == pid])
        writer.table(ws, 'Available season opportunity averages — separate scope', [r for r in opportunity_rows if str(r.get('PLAYER_ID')) == pid])
        writer.table(ws, 'Validated player rotation stints', [r for r in stints if str(r.get('player_id')) == pid])
        writer.table(ws, 'Actual historical team by season and phase', _group_summaries(history, ['season', 'season_type', 'team']))
        split_rows = []
        for row in regular:
            g = game_map.get(row['game_id'], {})
            venue = 'Home' if g.get('home') == row['team'] else 'Away' if g.get('away') == row['team'] else 'Unknown'
            days = rest.get((row['game_id'], row['team']))
            split_rows.append({**row, 'venue': venue, 'role': {1: 'Started', 0: 'Bench'}.get(row['started'], 'Unknown'), 'rest_bucket': 'Unknown' if days is None else str(min(days, 3)) + ('+' if days >= 3 else '')})
        for key in ('venue', 'role', 'rest_bucket', 'opponent'):
            writer.table(ws, 'Regular-season splits by ' + key, _group_summaries(split_rows, ['season', 'season_type', key]), note='Missing game context stays unknown. Source retrieval now does not imply a role was known before that game.')
        writer.table(ws, 'Verified historical game logs — newest first, phase retained', sorted(history, key=lambda r: (r['date_et'], r['game_id']), reverse=True), GAME_FIELDS)
        writer.table(ws, 'Unmatched archival game rows — not added to verified totals', sorted(older_history, key=lambda r: str(r.get('Game date', '')), reverse=True))
        writer.table(ws, 'Uploaded career rest splits — separate source scope', [r for r in raw_rest if r.get('Player') == player['name']])
    upcoming = [g for g in games if g.get('status') == 'scheduled' and (g.get('home') == 'NYK' or g.get('away') == 'NYK')]
    next_dates = {}
    for game in upcoming:
        opponent = game['away'] if game['home'] == 'NYK' else game['home']
        next_dates[opponent] = min(next_dates.get(opponent, '9999'), game['date_et'])
    opponents = sorted((t for t in NBA_TEAMS if t != 'NYK'), key=lambda t: (next_dates.get(t, '9999'), t))
    for team in opponents:
        ws = writer.sheet('NYK_vs_' + team, 'Knicks vs ' + team + ' — exact named pairs, not complete opposing rosters')
        relevant = by_team.get(team, set())
        these_pairs = [p for p in scope['pairs'] if p['opponent_team'] == team]
        writer.table(ws, 'Matchup coverage and explicit names', [{'opponent': team, 'next_scheduled_meeting': next_dates.get(team), 'canonical_through': cutoff, 'selected_knicks': ', '.join(p['name'] for p in selected), 'named_opponents': ', '.join(names[p] for p in sorted(relevant)), 'pair_count': len(these_pairs), 'scope': 'Only each selected Knicks player’s named primary/alternate opponent; estimated assignments remain unconfirmed'}])
        writer.table(ws, 'Specific named matchup pairs — primary and alternate only', [r for r in pair_rows if r['opponent_team'] == team])
        knicks = [r for r in rows if r['player_id'] in selected_ids and r.get('team') == 'NYK' and r.get('opponent') == team]
        opposing = [r for r in rows if r['player_id'] in relevant and r.get('team') == team and r.get('opponent') == 'NYK']
        writer.table(ws, 'Whole-game head-to-head summaries — Knicks first', _group_summaries(knicks, ['player_id', 'name', 'season', 'season_type'], side='Selected Knicks') + _group_summaries(opposing, ['player_id', 'name', 'season', 'season_type'], side='Named matchup opponent'), note='These are full-game totals against the team, not production while guarded by a specific player.')
        writer.table(ws, 'Broader regular-season baselines — named players only', [{'side': 'Selected Knicks', **s} for s in regular_annual if s['player_id'] in selected_ids] + [{'side': 'Named opponent', 'comparison_team_snapshot': team, **s} for s in regular_annual if s['player_id'] in relevant], note='Baseline history preserves actual teams. The comparison-team snapshot is not a retrospective roster or actual guarding assignment.')
        writer.table(ws, 'Selected Knicks full-game history — newest first', sorted(knicks, key=lambda r: (r['date_et'], r['game_id'], r['player_id']), reverse=True), GAME_FIELDS)
        writer.table(ws, 'Named opponent games against NYK — newest first', sorted(opposing, key=lambda r: (r['date_et'], r['game_id'], r['player_id']), reverse=True), GAME_FIELDS)
        writer.table(ws, 'Measured game-level exact pair history', [r for r in measured if r.get('opponent_team') == team])
        writer.table(ws, 'Game opportunities for these actual fixtures', [r for r in tracking_games + passing_games if r.get('game_id') in {g['game_id'] for g in knicks + opposing}])
        writer.table(ws, 'Validated scoped rotation stints in these fixtures', [r for r in stints if r.get('game_id') in {g['game_id'] for g in knicks + opposing}])
        writer.table(ws, 'Season opportunity means — separate aggregate', [r for r in opportunity_rows if str(r.get('PLAYER_ID')) in selected_ids | relevant])
        writer.table(ws, 'Uploaded measured defender summaries — regular season', [r for r in defenders if r.get('Opponent team') == team and r.get('Phase') == 'Regular season'])
        writer.table(ws, 'Uploaded measured defender summaries — overlapping all phases', [r for r in defenders if r.get('Opponent team') == team and r.get('Phase') != 'Regular season'], note='All(RS+PO) overlaps regular season; never add these phases together.')
        writer.table(ws, 'Estimated assignments — explicitly unconfirmed', these_pairs)
        writer.table(ws, 'Position defense — dated team/position definition', [r for r in position_defense if r.get('Team') in ('NYK', team)])
    ws = writer.sheet('Data_PlayerSeason', 'Selected-player and named-matchup season summaries; phases separate')
    writer.table(ws, 'Derived season summaries — fixed player scope and phase', annual)
    writer.table(ws, 'Independent career annual controls — selected Knicks', career_controls)
    writer.table(ws, 'Uploaded scoped summaries — separate dated source', raw_stats)
    ws = writer.sheet('Data_PlayerGames', 'Unique selected career games and named opponent games against NYK')
    writer.table(ws, 'Scoped verified player games — chronological', rows, GAME_FIELDS, 'One observation per game_id/player_id. Selected Knicks have general career history; opposing individual logs are actual NYK fixtures only. Canonical observations take precedence over overlap.')
    ws = writer.sheet('Data_TeamGames', 'Official team-game pace and defense context; no individual roster expansion')
    for sheet, dataset, measure in (('Team_games_base_2025', 'team_games_base', 'Base'), ('Team_games_advanced_2025', 'team_games_advanced', 'Advanced')):
        if ('2025-26', measure) not in team_scopes:
            writer.table(ws, sheet, _annotate(audit.get(sheet, []), dataset, sources, 'Official team-game totals or published pace/ratings'))
    for table in sorted(team_tables, key=lambda t: (str(t.get('season')), str(t.get('measure')))):
        writer.table(ws, 'Official team games ' + str(table.get('season')) + ' ' + str(table.get('measure')), _supplemental_rows(table))
    writer.table(ws, 'Team season advanced — separate aggregate', _annotate(audit.get('Team_season_advanced_2025', []), 'team_season_advanced', sources, 'Season aggregate, not pregame feature'))
    ws = writer.sheet('Data_GameContext', 'Available schedule and rest context for scoped games')
    writer.table(ws, 'Known game context — actual teams and Eastern times', contexts)
    writer.table(ws, 'Verified observed game-level starts — newly retrieved', observed_starts, note='Game starts remain separate from original box-stat known_at. Annual GS totals never fill per-game starting status.')
    writer.table(ws, 'Archived sample starting positions', _annotate([r for r in audit.get('Sample_game_box', []) if allowed_person(r)], 'sample_game_box', sources, 'Published game starting positions'))
    ws = writer.sheet('Data_TrackingGames', 'Verified scoped game-level opportunities; missing differs from zero')
    writer.table(ws, 'Verified game tracking — touches and rebound chances', tracking_games)
    writer.table(ws, 'Verified single-game passing — potential assists', passing_games, note='Returned source team labels remain separate from actual historical teams. Unique player/date joins require assist and published-minute agreement.')
    writer.table(ws, 'Archived one-date passing — scoped identities', _annotate([r for r in audit.get('Passing_one_date_2025', []) if allowed_person(r)], 'passing_single_date_check', sources, 'GP=1 source observations; not added again to current verified tables'))
    writer.table(ws, 'Archived sample game tracking — scoped identities', _annotate([r for r in audit.get('Sample_game_tracking', []) if allowed_person(r)], 'sample_game_tracking', sources, 'Dated source reference; overlapping samples not additional observations'))
    ws = writer.sheet('Data_TrackingSeason', 'Scoped season opportunity means; not individual games')
    for season, measure in sorted({(r.get('season', r.get('coverage_season')), r.get('reference_dataset')) for r in opportunity_rows}, key=lambda t: (str(t[0]), str(t[1])), reverse=True):
        group = [r for r in opportunity_rows if (r.get('season', r.get('coverage_season')), r.get('reference_dataset')) == (season, measure)]
        writer.table(ws, 'Season tracking ' + str(season) + ' ' + str(measure).split(':')[-1], group)
    ws = writer.sheet('Data_LineupsSeason', 'Knicks lineup combinations and named-player opposing team context only')
    lineup_data = audit.get('Lineup_season_combinations', [])
    def lineup_ids(row):
        return set(re.findall(r'\d+', str(row.get('GROUP_ID', ''))))
    nyk_lineups = [r for r in lineup_data if r.get('TEAM_ABBREVIATION') == 'NYK' and lineup_ids(r) & selected_ids]
    lineup_note = (f"The archived audit preserves {len(lineup_data)} source combinations from per-team requests. This view contains selected-player Knicks combinations; team context is not an individual absence effect."
                   if lineup_data else 'No archived season lineup combinations were supplied. Missing combinations remain unavailable; no private workbook references are assumed.')
    writer.table(ws, 'Knicks season lineup combinations involving selected players', _annotate(nyk_lineups, 'season_five_player_lineups', sources, 'Five-player team combinations; not individual absence effects'), note=lineup_note)
    lineup_summaries = []
    for team, ids in by_team.items():
        team_rows = [r for r in lineup_data if r.get('TEAM_ABBREVIATION') == team]
        for pid in sorted(ids):
            group = [r for r in team_rows if pid in lineup_ids(r)]
            values = [r.get('MIN') for r in group]
            lineup_summaries.append({'opponent_team': team, 'named_matchup_player_id': pid, 'name': names.get(pid), 'lineup_combinations': len(group), 'published_combination_minutes': sum(values) if values and all(isinstance(v, (int, float)) for v in values) else None, 'scope': 'Team combination coverage involving this named player; not their individual production or an absence-effect split'})
    writer.table(ws, 'Named opponent lineup coverage — no full roster rows', lineup_summaries)
    ws = writer.sheet('Data_RotationStints', 'Only validated scoped measured player stints and team intervals')
    writer.table(ws, 'Validated measured player rotation stints', stints)
    writer.table(ws, 'Validated five-player team intervals — scoped participants', lineup_intervals, note='Full unit size is validated internally. Other roster identities are not expanded into individual player pages or absence-effect estimates.')
    writer.table(ws, 'Rotation verification and unavailable games', rotations.get('checks', []) + rotations.get('errors', []))
    ws = writer.sheet('Data_DefenderPairs', 'Exact named defender pairs; whole-game and partial-possession scopes separate')
    writer.table(ws, 'Verified game-level primary or alternate named pairs', measured)
    writer.table(ws, 'Uploaded measured pair summaries — original phases preserved', defenders)
    writer.table(ws, 'Original unavailable regular-season pairs — dated source claims', unavailable_original, note='New measured game samples do not retroactively repair a missing career source summary. Missing original pairs remain distinguishable from zero.')
    ws = writer.sheet('Data_Shots', 'Verified scoped shot events; no invented defender-distance field')
    writer.table(ws, 'Verified shot events — selected Knicks and named matchups only', shots)
    current_shot_keys = {(str(r.get('game_id')), str(r.get('game_event_id')), str(r.get('player_id'))) for r in shots}
    archived_shots = [r for r in audit.get('Sample_shot_events', []) if allowed_person(r) and (str(r.get('GAME_ID')), str(r.get('GAME_EVENT_ID')), str(r.get('PLAYER_ID'))) not in current_shot_keys]
    writer.table(ws, 'Nonoverlapping archived sample shots', _annotate(archived_shots, 'sample_game_shots', sources, 'One official field-goal attempt per event; no defender-distance measurement'))
    ws = writer.sheet('Sources', 'Originals and retrospective source URLs/times for selected and paired views')
    writer.table(ws, 'Original source registry — provenance archive', _raw(store, 'Source Registry'))
    writer.table(ws, 'Historical and supplemental source catalog', sources)
    writer.table(ws, 'Original preserved files', store.rows('SELECT * FROM files ORDER BY imported_at'))
    writer.table(ws, 'Audit archive checksum', [audit_meta] if audit_meta else [])
    ws = writer.sheet('Data_Dictionary', 'Field meanings, source units, scope and historical knowledge limits')
    dictionary = [
        {'field': 'game_id + player_id', 'definition': 'Unique verified game observation. Canonical overlap takes precedence; selected career and playoff references do not mutate the model database.', 'units': 'Text IDs; leading zeros retained'},
        {'field': 'season_type', 'definition': 'Regular Season and Playoffs remain separate in annual, current, recent and career summaries.', 'units': 'Never combine overlapping phases or career/source aggregates'},
        {'field': 'team / opponent', 'definition': 'Actual game affiliations; matchup assignment teams are dated comparison scope only.', 'units': 'NBA historical abbreviation'},
        {'field': 'minutes / min_seconds / source_minutes / *_exact_text', 'definition': 'Canonical precision retained; supplementary published clock and raw numeric minutes stay separate. Text companions preserve numeric strings through Excel rounding.', 'units': 'Minutes / seconds / original published value / exact text representation'},
        {'field': 'started / started_source / started_known_at', 'definition': 'Observed start/bench role and separate source retrieval time. Annual GS does not establish a per-game start.', 'units': '1=start;0=bench; blank=unknown'},
        {'field': 'appeared', 'definition': 'Nonappearance is distinct from a played game with zero production.', 'units': '1=played;0=listed nonappearance'},
        {'field': 'specific named matchup pair', 'definition': 'Selected Knicks player plus their explicit primary/alternate opponent. Estimated assignments are not measured future guarding.', 'units': 'Exact stable identity pair and opponent team'},
        {'field': 'full-game / partialPossessions', 'definition': 'Full-game team-opponent production and measured partial-possession guarding are separate observations.', 'units': 'Played game / fractional tracked possessions'},
        {'field': 'season opportunity / game opportunity', 'definition': 'Season per-game means are not individual game observations; GP/scope/cutoff remain attached.', 'units': 'Published source units'},
        {'field': 'lineup combinations / validated stints', 'definition': 'Team context is not individual production with a teammate absent.', 'units': 'Published team combinations / validated source seconds'},
        {'field': 'retrieved_at / known_at', 'definition': 'New source retrieval is retrospective knowledge, not evidence it was available before an older game.', 'units': 'Unambiguous timestamps where supplied'},
    ]
    writer.table(ws, 'Scope and field dictionary', dictionary)
    writer.table(ws, 'Original uploaded data guide — dated semantics retained', _raw(store, 'Data Guide'))
    ws = writer.sheet('Data_Checks', 'Canonical checks and scoped supplemental verification/quarantines')
    writer.table(ws, 'Canonical validation — underlying league database', canonical_validation['checks'])
    writer.table(ws, 'Original archived data checks — dated source claims', _raw(store, 'Data Checks'))
    writer.table(ws, 'Scoped merge/source conflicts — source values retained', merge_conflicts)
    writer.table(ws, 'Independent career controls and quarantined packets', [t.get('checks', {}) | {'key': t.get('key')} for t in career.get('tables', [])] + career.get('quarantined_entries', []))
    writer.table(ws, 'Opportunity and rotation verification', opportunities.get('checks', []) + rotations.get('checks', []))
    writer.table(ws, 'Measured pair and shot-event source verification', matchups.get('coverage', []))
    writer.table(ws, 'Source failures and unresolved identities', errors + scope['unresolved'])
    writer.table(ws, 'Lineup minute checks — original discrepancies retained', audit.get('Lineup_minute_checks', []))
    writer.table(ws, 'Tracking team-label conflicts — scoped identities', [r for r in audit.get('Tracking_team_label_checks', []) if str(r.get('player_id')) in all_ids])
    ws = writer.sheet('Missing_Data', 'Remaining unavailable observations; no fabricated fields or old pregame snapshots')
    writer.table(ws, 'Remaining historical gaps and cautions', warnings)
    ws = writer.sheet('Update_Log', 'Source refresh history and current scoped workbook build')
    writer.table(ws, 'Current workbook build', [{'generated_at': created, 'canonical_through': cutoff, 'last_successful_refresh': last_success['finished_at'] if last_success else None, 'selection': [p['id'] for p in selected], 'exported_unique_games': len(rows), 'scope': scope['meaning'], 'historical_snapshot_limit': 'New retrievals do not create historical pregame knowledge'}])
    writer.table(ws, 'Operational source freshness', store.rows('SELECT * FROM freshness ORDER BY kind'))
    writer.table(ws, 'Update runs', store.rows('SELECT * FROM update_runs ORDER BY started_at'))
    audits = _scoped_correction_audit(store.rows('SELECT * FROM audit ORDER BY id'), rows)
    writer.table(ws, 'Scoped correction audit — original old and new values retained', audits)
    ws = writer.sheet('Contents', 'Linked contents — selected Knicks, exact named opponent pairs, supporting data')
    writer.line(ws, [coverage_note])
    writer.line(ws, [f"First{len(selected)} player pages, next29 opponent pages, then{len(SUPPORTING_SHEETS)} support sheets. Opponent logs are NYK fixtures only; broader baselines are named-player summaries."])
    writer.line(ws, ['Sheet', 'Purpose'], bold=True)
    for entry in writer.contents:
        cell = WriteOnlyCell(ws, value=entry['sheet'])
        cell.data_type = 's'
        cell.hyperlink = "#'" + entry['sheet'].replace("'", "''") + "'!A1"
        cell.font = Font(color='0563C1', underline='single')
        ws.append([cell, entry['purpose']])
        writer.positions[ws.title] += 1
    writer.table(ws, 'Table row counts and ranges', writer.tables[:])
    expected_names = list(writer.positions)
    temporary = destination.parent / ('.' + destination.stem + '-' + uuid.uuid4().hex + '.xlsx')
    try:
        writer.book.save(temporary)
        check = load_workbook(temporary, read_only=True, data_only=False)
        try:
            if check.sheetnames != expected_names or len(check.sheetnames) != len(selected) + 29 + len(SUPPORTING_SHEETS):
                raise ValueError('Organized workbook sheet ordering/count check failed')
            if check.sheetnames[len(selected) + 29:] != SUPPORTING_SHEETS:
                raise ValueError('Supporting sheet order check failed')
        finally:
            check.close()
        os.replace(temporary, destination)
    except BaseException:
        temporary.unlink(missing_ok=True)
        raise
    report = {'status': 'passed_with_limitations', 'workbook': str(destination), 'generated_at': created, 'canonical_through': cutoff,
              'last_successful_refresh': last_success['finished_at'] if last_success else None, 'selected_players': selected,
              'sheets': expected_names, 'canonical_player_games': len(canonical), 'canonical_unique_player_games': len(canonical_by_key),
              'exported_unique_player_games': len(rows), 'selected_verified_career_games': len(played),
              'original_unavailable_pairs': len(unavailable_original), 'original_pairs_observed_in_new_regular_season_window': recovered_original_pairs, 'selected_unknown_starting_roles': unknown_starts, 'canonical_selected_unknown_roles_before': original_unknown_starts, 'canonical_selected_unknown_roles_after': remaining_canonical_unknown_starts,
              'scope': scope, 'reference': audit_meta, 'tables': writer.tables, 'validation': canonical_validation,
              'supplemental_context': {k: v for k, v in supplemental.items() if k != 'tables'},
              'historical_reference_coverage': {kind: {k: reference.get(k) for k in ('status', 'coverage', 'summary', 'errors') if k in reference} for kind, reference in references.items()},
              'unavailable_and_cautions': warnings, 'merge_conflicts': merge_conflicts, 'sha256': digest(destination.read_bytes())}
    (destination.parent / 'organized_validation.json').write_text(json.dumps(report, indent=2, default=str))
    with (destination.parent / 'organized_missing.csv').open('w', newline='', encoding='utf-8') as f:
        output = csv.DictWriter(f, fieldnames=['code', 'availability', 'detail'])
        output.writeheader()
        output.writerows(warnings)
    change_log = ('# Scoped historical workbook update\n\nGenerated ' + created + '. Canonical coverage through ' + str(cutoff) + '.\n\n'
                  '- Created ' + str(len(selected)) + ' explicitly selected player pages,29 exact named-matchup opponent pages and16 supporting sheets.\n'
                  '- Preserved league history in the underlying database while exporting selected-player career history and named opponent games against NYK only. No full opponent roster expansion.\n'
                  '- Added verified available career box scores/playoffs, game opportunities, starts, validated rotations, exact guarding pairs and shot events, each with its own coverage and current retrieval timestamps.\n'
                  '- Kept regular season and playoffs separate, deduplicated overlapping source observations, and retained canonical minute precision, missing values, original files and source conflicts.\n'
                  '- Kept actual full-game H2H production separate from measured partial-possession guarding and estimated future assignments.\n'
                  '- Historical pregame vintages remain unavailable where not archived before the game; no game-by-game injury-history tables or invented statistics were added.\n')
    (destination.parent / 'organized_changelog.md').write_text(change_log)
    return report

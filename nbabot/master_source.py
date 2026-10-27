"""Private, offline ingestion of one user-supplied master spreadsheet.

Uploaded source timestamps and upload time are different facts.  This module
never follows links, contacts a provider, or marks an old observation current.
The active metadata is committed with the derived records in SQLite; the JSON
pointer is only a recoverable convenience copy.
"""
from __future__ import annotations

import csv
import fcntl
import json
import math
import os
import posixpath
import re
import shutil
import tempfile
import uuid
import zipfile
from contextlib import contextmanager
from datetime import datetime
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

from .core import ET, digest, dumps, now, stamp
from .ingest import (COUNTS, PG_FIELDS, date_only, duration, key, normalize_game,
                     source_time, upsert_game, upsert_player_game, validate)

MAX_FILE_BYTES=32*1024*1024
MAX_UNPACKED_BYTES=1024*1024*1024
MAX_ZIP_MEMBERS=5000
MAIN_NS="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
REL_NS="http://schemas.openxmlformats.org/package/2006/relationships"
OFFICE_REL="http://schemas.openxmlformats.org/officeDocument/2006/relationships"
CANONICAL_SHEETS={"dataplayergames","cleanstats","nbaplayergames","playergames"}


def _require_mode(store):
    if store.config.get("data_source_mode")!="spreadsheet":
        raise ValueError("Master replacement requires data_source_mode='spreadsheet'")


@contextmanager
def _maintenance_lock(store,lock_held=False):
    if lock_held:
        yield
        return
    with (store.root/".refresh.lock").open("a") as lock:
        try:
            fcntl.flock(lock,fcntl.LOCK_EX|fcntl.LOCK_NB)
        except BlockingIOError as exc:
            raise ValueError("Another master import or refresh is running") from exc
        yield


def _atomic_json(path,payload):
    temp=path.parent/("."+path.name+"-"+uuid.uuid4().hex)
    try:
        temp.write_text(dumps(payload))
        os.replace(temp,path)
    finally:
        temp.unlink(missing_ok=True)


def master_status(store):
    row=store.one("SELECT payload FROM service_status WHERE name='master_source'")
    if not row:
        return {"status":"waiting_for_master","file_hash":None,"relative_path":None,
            "original_name":None,"uploaded_at":None,"coverage":None,"last_local_refresh":None,
            "note":"Upload a master XLSX. No outside data source is enabled."}
    result=json.loads(row["payload"])
    refresh=store.one("SELECT payload FROM service_status WHERE name='master_refresh'")
    result["last_local_refresh"]=json.loads(refresh["payload"]) if refresh else None
    return result


def master_path(store):
    """Resolve only the accepted original, never arbitrary paths in a workbook."""
    status=master_status(store)
    relative=status.get("relative_path")
    if not relative:
        return None
    name=PurePosixPath(relative)
    if name.is_absolute() or len(name.parts)!=2 or name.parts[0]!="originals" or ".." in name.parts:
        raise ValueError("Invalid active master path")
    path=store.root/Path(*name.parts)
    if path.is_symlink() or not path.resolve().is_relative_to((store.root/"originals").resolve()):
        raise ValueError("Active master must be an archived original file")
    return path


def _check_archive(path):
    if path.suffix.lower()!=".xlsx":
        raise ValueError("Upload an .xlsx workbook")
    if not path.is_file() or not 0<path.stat().st_size<=MAX_FILE_BYTES:
        raise ValueError("Master XLSX must be nonempty and at most32MiB")
    try:
        with zipfile.ZipFile(path) as archive:
            infos=archive.infolist()
            if len(infos)>MAX_ZIP_MEMBERS or sum(i.file_size for i in infos)>MAX_UNPACKED_BYTES:
                raise ValueError("Workbook exceeds safe expanded-size limits")
            names=set()
            for info in infos:
                name=PurePosixPath(info.filename)
                if name.is_absolute() or ".." in name.parts or "\\" in info.filename or info.filename in names:
                    raise ValueError("Unsafe or duplicate XLSX archive member")
                if info.flag_bits&1:
                    raise ValueError("Encrypted workbooks must be unlocked before import")
                if "vbaproject" in info.filename.lower() or info.filename.lower().endswith(".bin"):
                    raise ValueError("Macro/binary workbook content is unsupported")
                names.add(info.filename)
            if not {"xl/workbook.xml","[Content_Types].xml"}.issubset(names):
                raise ValueError("File is not an Excel workbook")
            if archive.testzip() is not None:
                raise ValueError("Workbook archive checksum failed")
    except zipfile.BadZipFile as exc:
        raise ValueError("File is not a readable XLSX archive") from exc


def _tables(path):
    """Read tiny table-relationship files, not large worksheet XML documents."""
    output={}
    with zipfile.ZipFile(path) as z:
        workbook=ElementTree.fromstring(z.read("xl/workbook.xml"))
        rels=ElementTree.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        targets={r.attrib["Id"]:r.attrib["Target"] for r in rels if r.attrib.get("TargetMode")!="External"}
        for sheet in workbook.findall("{"+MAIN_NS+"}sheets/{"+MAIN_NS+"}sheet"):
            target=targets.get(sheet.attrib.get("{"+OFFICE_REL+"}id"))
            if not target:continue
            worksheet=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl",target))
            relation=posixpath.join(posixpath.dirname(worksheet),"_rels",posixpath.basename(worksheet)+".rels")
            specs=[]
            if relation in z.namelist():
                for rel in ElementTree.fromstring(z.read(relation)):
                    if not rel.attrib.get("Type","").endswith("/table") or rel.attrib.get("TargetMode")=="External":continue
                    t=rel.attrib["Target"]
                    member=t.lstrip("/") if t.startswith("/") else posixpath.normpath(posixpath.join(posixpath.dirname(worksheet),t))
                    if not member.startswith("xl/tables/"):raise ValueError("Invalid XLSX table relationship")
                    table=ElementTree.fromstring(z.read(member))
                    specs.append({"name":table.attrib.get("displayName",member),"range":range_boundaries(table.attrib["ref"]),
                        "headers":[c.attrib["name"] for c in table.findall("{"+MAIN_NS+"}tableColumns/{"+MAIN_NS+"}tableColumn")]})
            output[sheet.attrib["name"]]=sorted(specs,key=lambda s:(s["range"][1],s["range"][0]))
    return output


def _records(path):
    table_specs=_tables(path)
    book=load_workbook(path,read_only=True,data_only=True,keep_links=False)
    try:
        for sheet in book:
            specs=table_specs.get(sheet.title,[])
            if specs:
                active=[]; pending=iter(specs); upcoming=next(pending,None)
                for number,values in enumerate(sheet.iter_rows(values_only=True),1):
                    while upcoming and upcoming["range"][1]<=number:
                        active.append(upcoming);upcoming=next(pending,None)
                    for spec in active:
                        left,top,right,bottom=spec["range"]
                        if number==top:
                            actual=list(values[left-1:right])
                            if [str(v) for v in actual]!=spec["headers"]:
                                raise ValueError("Excel table headers disagree with cells on "+sheet.title)
                        elif top<number<=bottom:
                            row=dict(zip(spec["headers"],values[left-1:right]))
                            if any(v is not None for v in row.values()):
                                yield sheet.title,spec["name"],number,row
                    active=[s for s in active if s["range"][3]>number]
                    if upcoming is None and not active:break
            else:
                headers=None;data_seen=False;ended=False
                for number,values in enumerate(sheet.iter_rows(values_only=True),1):
                    nonempty=[v for v in values if v is not None]
                    if headers is None:
                        if len(nonempty)>=2 and all(isinstance(v,str) for v in nonempty):
                            headers=[str(v) if v is not None else "unnamed_"+str(i) for i,v in enumerate(values)]
                        continue
                    # Some spreadsheet editors remove Excel Table metadata.
                    # These canonical/context sheets still have exactly one
                    # first dataset; a later role/header section is not a game.
                    if key(sheet.title) in CANONICAL_SHEETS|{"datagamecontext"}:
                        if not nonempty and data_seen:ended=True
                        if ended:
                            if nonempty:yield sheet.title,"uninterpreted_after_first_table",number,{"__raw_cells":list(values)}
                            continue
                    row=dict(zip(headers,values))
                    if any(v is not None for v in row.values()):
                        data_seen=True
                        yield sheet.title,"flat",number,row
    finally:
        book.close()


def _numeric_exact(row,field,tolerance):
    numeric=row.get(field)
    exact=row.get(field+"_exact_text")
    if exact in (None,""):return numeric
    try:
        e=float(exact)
        if not math.isfinite(e):raise ValueError
        if numeric not in (None,"") and (not math.isfinite(float(numeric)) or abs(float(numeric)-e)>tolerance):
            raise ValueError
        return e
    except (TypeError,ValueError) as exc:
        raise ValueError("Exact and numeric "+field+" disagree") from exc


def _normalized(row,uploaded_at,timestamps,file_hash,identities):
    raw=dict(row)
    canonical=any(key(k)=="gameid" for k in raw) and "game_id" in raw
    if canonical:
        raw["PLAYER_NAME"]=raw.get("name") or identities.get(str(raw.get("player_id")))
        raw["Minutes (source precision)"]=_numeric_exact(raw,"minutes",1e-9)
        raw["Minutes in seconds"]=_numeric_exact(raw,"min_seconds",1e-6)
        raw["Original reported MIN"]=raw.get("source_minutes")
    source=raw.get("Source ID",raw.get("source"))
    known=raw.get("known_at") or raw.get("retrieved_at") or timestamps.get(source)
    valid_known=None
    if known:
        try:valid_known=stamp(known)
        except (ValueError,TypeError):pass
    raw["known_at"]=valid_known or uploaded_at
    raw["Master workbook SHA256"]=file_hash
    raw["Master timestamp basis"]=("Provided source timestamp; not a new online verification" if valid_known else
        "First observed in uploaded master; original retrieval time unavailable")
    p=normalize_game(raw,uploaded_at)
    if not re.fullmatch(r"\d{10}",p["game_id"]) or not p["game_id"].startswith(("002","004")):
        raise ValueError("Canonical player games need a ten-digit regular-season or playoff NBA game ID")
    if not re.fullmatch(r"\d+",p["player_id"]):raise ValueError("Canonical observations need a stable numeric NBA player ID")
    expected="Playoffs" if p["game_id"].startswith("004") else "Regular Season"
    phase=str(row.get("season_type") or expected)
    if phase not in (expected,"Regular season" if expected=="Regular Season" else expected):
        raise ValueError("Game phase disagrees with stable game ID")
    if p["date_et"]>datetime.now(ET).date().isoformat():raise ValueError("Future player-game observation")
    if re.fullmatch(r"\d{4}-\d{2}",p["season"]):
        start=int(p["season"][:4]);date=datetime.fromisoformat(p["date_et"])
        if int(p["season"][-2:])!=(start+1)%100 or date.year not in (start,start+1) or (date.year==start and date.month<7):
            raise ValueError("Player-game date disagrees with supplied season")
        if p["game_id"][3:5]!=p["season"][2:4]:raise ValueError("Game ID disagrees with supplied season")
    if p["appeared"] is None:raise ValueError("Player games must distinguish played, listed nonappearance and missing observations")
    if p["appeared"]==0:
        # Published placeholder zeros remain in payload; they are not played production.
        for field in ["minutes","min_seconds",*COUNTS]:p[field]=None
        p["started"]=None
        p["hash"]=digest({k:v for k,v in p.items() if k not in ("known_at","source","payload","hash","name","team_id","opponent_id")})
    p["phase"]=expected
    return p


def _same(a,b):
    for name in PG_FIELDS:
        if name in ("known_at","source","payload","hash"):continue
        left,right=a.get(name),b.get(name)
        if name in ("minutes","min_seconds") and left is not None and right is not None:
            if not math.isclose(left,right,rel_tol=0,abs_tol=1e-6):return False
        elif left!=right:return False
    return True


def _schedule(row,uploaded_at,season):
    if "game_id" in row or "id" in row:
        gid=row.get("game_id",row.get("id"))
        if not gid:return None
        known=row.get("known_at") or uploaded_at
        try:known=stamp(known)
        except (TypeError,ValueError):known=uploaded_at
        return {"id":str(gid),"date_et":date_only(row["date_et"]),"tipoff":row.get("tipoff"),
            "home":row.get("home"),"away":row.get("away"),"status":str(row.get("status") or "unknown").lower(),
            "season":row.get("season") or season,"known_at":known,"source":row.get("source") or "uploaded master",
            "metadata":dumps(row)}
    if row.get("Date (ET)") and row.get("ESPN game ID"):
        date=date_only(row["Date (ET)"])
        clock=datetime.strptime(str(row["Tipoff (ET)"]),"%I:%M %p").time()
        tip=datetime.combine(datetime.fromisoformat(date).date(),clock,ET)
        opponent=row["Opponent"]
        home,away=("NYK",opponent) if row.get("Home / away")=="Home" else (opponent,"NYK")
        return {"id":"espn:"+str(row["ESPN game ID"]),"date_et":date,"tipoff":tip.isoformat(),"home":home,"away":away,
            "status":str(row.get("Status") or "unknown").lower(),"season":season,"known_at":uploaded_at,
            "source":row.get("Game source") or "uploaded master","metadata":dumps(row)}
    return None


def _archive_original(store,path,file_hash):
    root=store.root/"originals";root.mkdir(exist_ok=True)
    target=root/(file_hash+"_Master.xlsx")
    if target.exists():
        if digest(target.read_bytes())!=file_hash:raise ValueError("Archived master checksum mismatch")
    else:
        temp=root/(".master-"+uuid.uuid4().hex)
        try:
            shutil.copyfile(path,temp)
            if digest(temp.read_bytes())!=file_hash:raise ValueError("Master copy checksum mismatch")
            os.replace(temp,target)
        finally:temp.unlink(missing_ok=True)
    return target


def import_master(store,file_path,original_name=None,lock_held=False):
    _require_mode(store)
    with _maintenance_lock(store,lock_held):
        path=Path(file_path)
        _check_archive(path)
        file_hash=digest(path.read_bytes())
        active=master_status(store)
        if active.get("file_hash")==file_hash:
            accepted=master_path(store)
            if not accepted or not accepted.is_file() or digest(accepted.read_bytes())!=file_hash:
                raise ValueError("Accepted master is missing or changed; restore the original")
            return {"status":"unchanged","master":active,"regular_games":active["coverage"]["player_games"],
                "playoff_games":active["coverage"]["playoff_games"],"duplicates_removed":active.get("duplicates_removed",0)}
        uploaded_at=now();name=Path(original_name or path.name).name
        if not name.lower().endswith(".xlsx") or len(name)>255 or any(c in name for c in "\x00\r\n"):
            raise ValueError("Master filename must be a normal .xlsx basename")
        canonical_raw=[];contexts=[];timestamps={};identities={};team_ids={};team_conflicts=[];pairs=[];inventory={};raw_count=0;ordinals={}
        with tempfile.TemporaryDirectory(prefix="nbabot-master-") as directory:
            rawfile=Path(directory)/"raw.jsonl"
            with rawfile.open("w",encoding="utf-8") as archive:
                for sheet,table,number,row in _records(path):
                    clean={k:v for k,v in row.items() if k is not None}
                    ordinals[sheet]=ordinals.get(sheet,0)+1
                    archive.write(dumps({"sheet":sheet,"table":table,"number":ordinals[sheet],"excel_row":number,"payload":clean})+"\n")
                    raw_count+=1;inventory[sheet]=inventory.get(sheet,0)+1
                    normalized={key(k):v for k,v in clean.items()}
                    if key(sheet) in CANONICAL_SHEETS:
                        canonical_raw.append(clean)
                    if key(sheet)=="datagamecontext" and {"gameid","dateet","status"}.issubset(normalized) and "playerid" not in normalized:
                        contexts.append(clean)
                    if key(sheet)=="schedule" and "dateet" in normalized:
                        contexts.append(clean)
                    sid=clean.get("Source ID")
                    if sid:
                        at=source_time(clean.get("Retrieved at (UTC)")) or source_time(clean.get("Extracted at (UTC)"))
                        if at:timestamps[sid]=at
                    pid=clean.get("NBA player ID",clean.get("PLAYER_ID"))
                    pname=clean.get("Player",clean.get("PLAYER_NAME"))
                    if key(sheet)=="players":pid,pname=clean.get("id"),clean.get("name")
                    if pid is not None and pname:
                        stable=str(int(pid)) if isinstance(pid,(int,float)) and int(pid)==pid else str(pid).strip()
                        identities[stable]=str(pname)
                    tid=clean.get("NBA team ID",clean.get("TEAM_ID",clean.get("team_id")))
                    abbr=clean.get("TEAM_ABBREVIATION",clean.get("TEAM_ABBR",clean.get("Team",clean.get("team"))))
                    if tid is not None and abbr and re.fullmatch(r"161061\d{4}",str(tid)) and re.fullmatch(r"[A-Z]{2,4}",str(abbr)):
                        tid,abbr=str(tid),str(abbr)
                        if tid in team_ids and team_ids[tid]!=abbr:
                            team_conflicts.append({"team_id":tid,"retained_abbreviation":team_ids[tid],"other_supplied_abbreviation":abbr})
                        else:team_ids[tid]=abbr
                    if {"knicks_id","knicks_name","opponent_id","opponent_name","opponent_team","assignment"}.issubset(clean):
                        # The eight-column identity table is authoritative for
                        # estimated-rule scope, not a measured future assignment.
                        if len([k for k in clean if not k.startswith("unnamed")])<=10:
                            pairs.append(clean)
                            identities[str(clean["knicks_id"])]=str(clean["knicks_name"])
                            identities[str(clean["opponent_id"])]=str(clean["opponent_name"])
            if not canonical_raw:raise ValueError("No supported canonical game table found (Clean Stats or Data_PlayerGames)")
            games={};duplicates=0
            for row in canonical_raw:
                p=_normalized(row,uploaded_at,timestamps,file_hash,identities)
                previous=games.get((p["game_id"],p["player_id"]))
                if previous:
                    if not _same(previous,p):raise ValueError("Conflicting duplicate player-game records in master")
                    duplicates+=1
                else:games[(p["game_id"],p["player_id"])]=p
            regular=[p for p in games.values() if p["phase"]=="Regular Season"]
            playoff=[p for p in games.values() if p["phase"]=="Playoffs"]
            if not regular:raise ValueError("Master needs regular-season player observations for the baseline")
            schedules=[g for r in contexts if (g:=_schedule(r,uploaded_at,store.config["season"])) is not None]
            # Check identities/dates before modifying the live connection.
            dates={}
            teams={}
            for p in games.values():
                if p["game_id"] in dates and dates[p["game_id"]]!=p["date_et"]:raise ValueError("Game ID joins different dates")
                dates[p["game_id"]]=p["date_et"]
                if p.get("opponent"):
                    pair=frozenset([p["team"],p["opponent"]])
                    if len(pair)!=2:raise ValueError("Player team and opponent must differ")
                    if p["game_id"] in teams and teams[p["game_id"]]!=pair:raise ValueError("Game ID joins different actual teams")
                    teams[p["game_id"]]=pair
            aliases=[]
            for g in schedules:
                if g["id"].startswith("espn:") and g.get("home") and g.get("away"):
                    pair=frozenset([g["home"],g["away"]])
                    candidates=[gid for gid,date in dates.items() if date==g["date_et"] and teams.get(gid)==pair]
                    if len(candidates)>1:raise ValueError("Ambiguous schedule alias joins more than one canonical game")
                    if candidates:
                        aliases.append((g["id"],candidates[0]));g["id"]=candidates[0]
                if g["id"] in dates and dates[g["id"]]!=g["date_et"]:raise ValueError("Schedule and player-game dates conflict")
                if g.get("home") and g.get("away") and g["id"] in teams and frozenset([g["home"],g["away"]])!=teams[g["id"]]:
                    raise ValueError("Schedule and player-game teams conflict")
            original=_archive_original(store,path,file_hash)
            status={"status":"ready","file_hash":file_hash,"relative_path":original.relative_to(store.root).as_posix(),
                "original_name":name,"uploaded_at":uploaded_at,"coverage":{"first":min(p["date_et"] for p in games.values()),
                "last":max(p["date_et"] for p in games.values()),"player_games":len(regular),"playoff_games":len(playoff),
                "played_regular_games":sum(p["appeared"]==1 for p in regular)},
                "provided_source_times":sum("Provided source timestamp" in json.loads(p["payload"])["Master timestamp basis"] for p in games.values()),
                "raw_rows":raw_count,"duplicates_removed":duplicates,"schedule_aliases":len(aliases),"sheets":inventory,
                "supplied_team_ids":len(team_ids),"team_identity_conflicts":team_conflicts,
                "note":"Only this supplied master is used. Upload time does not refresh source observations; rosters, injuries and current offers are unknown unless explicitly modeled from supported supplied data."}
            with store.transaction():
                for table in ["player_games","stat_versions","projections","decisions","settlements","entries","offers","boards","rules","evaluations","model_runs",
                    "game_aliases","game_versions","games","players","teams","raw_rows","snapshots","freshness","checks","audit","files"]:
                    store.db.execute("DELETE FROM "+table)
                store.db.execute("DELETE FROM service_status WHERE name IN ('master_source','master_refresh')")
                store.db.execute("DELETE FROM jobs WHERE key LIKE 'pregame:%'")
                store.db.execute("INSERT INTO files VALUES(?,?,?,?,?)",(file_hash,name,str(original),uploaded_at,dumps(inventory)))
                for pid,pname in identities.items():store.db.execute("INSERT OR REPLACE INTO players VALUES(?,?)",(pid,pname))
                for tid,abbr in team_ids.items():
                    existing=store.one("SELECT id FROM teams WHERE abbr=?",(abbr,))
                    if existing and existing["id"]!=tid:
                        raise ValueError("Supplied team abbreviation joins multiple stable IDs")
                    store.db.execute("INSERT INTO teams VALUES(?,?)",(tid,abbr))
                for g in schedules:upsert_game(store,g)
                for alias,gid in aliases:store.db.execute("INSERT INTO game_aliases VALUES(?,?)",(alias,gid))
                for p in sorted(regular,key=lambda r:(r["date_et"],r["game_id"],r["player_id"])):
                    upsert_player_game(store,p)
                with rawfile.open(encoding="utf-8") as archive:
                    batch=[]
                    for line in archive:
                        row=json.loads(line)
                        payload={**row["payload"],"__master_table":row["table"],"__master_excel_row":row["excel_row"]}
                        batch.append((file_hash,row["sheet"],row["number"],dumps(payload)))
                        if len(batch)>=500:
                            store.db.executemany("INSERT INTO raw_rows VALUES(?,?,?,?)",batch);batch=[]
                    if batch:store.db.executemany("INSERT INTO raw_rows VALUES(?,?,?,?)",batch)
                for n,p in enumerate(playoff,1):
                    store.db.execute("INSERT INTO raw_rows VALUES(?,?,?,?)",(file_hash,"Master_Playoff_Games",n,dumps({**json.loads(p["payload"]),"season_type":"Playoffs"})))
                derived={}
                for pair in pairs:
                    k=(str(pair["knicks_id"]),str(pair["opponent_team"]))
                    r=derived.setdefault(k,{"Knicks player":pair["knicks_name"],"Opponent":pair["opponent_team"],"Assignment status":pair.get("assignment_status"),"Source ID":"master:"+file_hash})
                    field={"primary":"Suspected defender","alternate":"Alternate defender"}.get(str(pair["assignment"]).lower())
                    if field:
                        if field in r and r[field]!=pair["opponent_name"]:raise ValueError("Conflicting named matchup assignment in master")
                        r[field]=pair["opponent_name"]
                # Original Matchups rows are already retained; only organized
                # views require rebuilding their explicit identity-rule table.
                if derived and not store.one("SELECT 1 FROM raw_rows WHERE sheet='Matchups'"):
                    for n,r in enumerate(derived.values(),1):
                        store.db.execute("INSERT INTO raw_rows VALUES(?,?,?,?)",(file_hash,"Matchups",n,dumps(r)))
                report=validate(store)
                if report["status"]=="failed":
                    raise ValueError("Master validation failed: "+dumps([r for r in report["checks"] if r["severity"]=="error"][:10]))
                status["validation"]=report
                store.db.execute("INSERT INTO service_status VALUES('master_source',?,?)",(uploaded_at,dumps(status)))
            # This copy may be regenerated from committed DB metadata after a
            # crash. It is never authoritative for choosing an arbitrary file.
            try:_atomic_json(store.root/"master_source.json",status)
            except OSError:status["pointer_note"]="Advisory JSON could not be written; committed database metadata remains authoritative"
            return {"status":"success","master":status,"regular_games":len(regular),"playoff_games":len(playoff),"duplicates_removed":duplicates}


def refresh_master(store,lock_held=False):
    _require_mode(store)
    with _maintenance_lock(store,lock_held):
        status=master_status(store)
        if status["status"]=="waiting_for_master":return status
        path=master_path(store)
        if not path or not path.is_file() or digest(path.read_bytes())!=status["file_hash"]:
            raise ValueError("Accepted master is missing or has changed; upload and validate a new master")
        at=now();warnings=[];projections=[]
        try:
            from .model import project_upcoming
            projections=project_upcoming(store)
        except (ValueError,KeyError,TypeError) as exc:
            store.db.rollback()
            store.db.execute("UPDATE projections SET status='superseded' WHERE status='active'")
            store.db.execute("UPDATE decisions SET status='superseded' WHERE status='active'")
            store.db.commit()
            warnings.append({"code":"projections_withheld","detail":str(exc)})
        from .report import write_report
        try:
            output=write_report(store,store.root/"local_report")
        except (ValueError,KeyError,TypeError) as exc:
            output={"status":"unavailable","error_type":type(exc).__name__}
            warnings.append({"code":"report_withheld","detail":str(exc)})
        result={"status":"unchanged","master_unchanged":True,"recalc_done":True,"refreshed_at":at,
            "file_hash":status["file_hash"],"data_cutoff":status["coverage"]["last"],
            "projection_count":len(projections),"warnings":warnings,"report":output,
            "recommendations":"Withheld without current supplied offers, rules and verified assumptions",
            "note":"Recalculated only from the accepted master. No external data was fetched, no games were added and source freshness was not advanced."}
        store.db.execute("INSERT OR REPLACE INTO service_status VALUES('master_refresh',?,?)",(at,dumps(result)))
        store.db.commit()
        _atomic_json(store.root/"master_source.json",master_status(store))
        return result


def load_master_reference(store):
    """Optional reference access remains confined to records in the accepted file."""
    status=master_status(store)
    if status["status"]!="ready":return {"status":"waiting_for_master","tables":{},"playoff_games":[]}
    tables={}
    for row in store.rows("SELECT sheet,payload FROM raw_rows WHERE file_hash=? ORDER BY sheet,row_number",(status["file_hash"],)):
        tables.setdefault(row["sheet"],[]).append(json.loads(row["payload"]))
    return {"status":"supplied_master_only","file_hash":status["file_hash"],"tables":tables,
        "playoff_games":tables.get("Master_Playoff_Games",[]),"source_freshness":"Provided timestamps only; no new outside verification"}
